import io
import json
from datetime import date
from xml.sax.saxutils import escape as xml_escape

from django.utils.text import slugify
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.util import Inches as PptInches, Pt as PptPt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import BaseDocTemplate, Frame, Image, KeepTogether, ListFlowable, ListItem, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle

from .document_branding import BRAND, LOGO_PATH, briefing_sections, clean_text, proposal_sections


MIMES = {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation", "md": "text/markdown; charset=utf-8", "json": "application/json; charset=utf-8"}


def export_filename(opportunity, extension):
    return f"briefing-{slugify(opportunity.title) or f'oportunidade-{opportunity.pk}'}.{extension}"


def proposal_export_filename(opportunity, version, extension):
    project = slugify(opportunity.title) or f"oportunidade-{opportunity.pk}"
    return f"proposta-{project}-{version.lower()}.{extension}"


def public_payload(briefing):
    return {"title": briefing.opportunity.title, "client": briefing.opportunity.client_name, "updated_at": briefing.updated_at.isoformat(), "briefing": {s.key: list(s.items) if isinstance(s.value, tuple) else s.value for s in briefing_sections(briefing)}}


def _document_date(briefing):
    value = getattr(briefing, "updated_at", None)
    return (value.date() if value else date.today()).strftime("%d/%m/%Y")


def _metadata(briefing):
    result = [("Documento", "Briefing profissional"), ("Projeto", clean_text(briefing.opportunity.title))]
    if clean_text(briefing.opportunity.client_name): result.append(("Cliente", clean_text(briefing.opportunity.client_name)))
    result.append(("Data", _document_date(briefing)))
    return result


def as_markdown(briefing):
    lines = [f"# {BRAND.organization}", "", "## Briefing profissional", ""]
    for label, value in _metadata(briefing)[1:]: lines.extend((f"**{label}:** {value}", ""))
    for section in briefing_sections(briefing):
        lines.extend((f"## {section.title}", ""))
        lines.extend((f"- {item}" for item in section.items) if isinstance(section.value, tuple) else (section.value,))
        lines.append("")
    return "\n".join(lines).encode("utf-8")


def _pdf_styles():
    sample = getSampleStyleSheet()
    return {
        "body": ParagraphStyle("3DS Body", parent=sample["BodyText"], fontName="Helvetica", fontSize=10.2, leading=14, textColor=colors.HexColor("#1C1C1C"), spaceAfter=6),
        "section": ParagraphStyle("3DS Section", parent=sample["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, textColor=colors.HexColor("#0057B8"), spaceBefore=10, spaceAfter=7, keepWithNext=True),
        "small": ParagraphStyle("3DS Small", parent=sample["BodyText"], fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor("#667085")),
        "coverbrand": ParagraphStyle("3DS Cover Brand", parent=sample["BodyText"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=colors.HexColor("#0057B8"), alignment=1),
        "covertitle": ParagraphStyle("3DS Cover Title", parent=sample["Title"], fontName="Helvetica-Bold", fontSize=28, leading=33, textColor=colors.HexColor("#1C1C1C"), alignment=1, spaceAfter=10),
        "coverproject": ParagraphStyle("3DS Cover Project", parent=sample["Heading2"], fontName="Helvetica", fontSize=16, leading=21, textColor=colors.HexColor("#0057B8"), alignment=1),
    }


def _pdf_page(canvas, doc):
    if doc.page == 1: return
    canvas.saveState(); width, _ = A4
    canvas.setStrokeColor(colors.HexColor("#D8E4F2")); canvas.line(20*mm, 281*mm, width-20*mm, 281*mm); canvas.line(20*mm, 15*mm, width-20*mm, 15*mm)
    canvas.setFont("Helvetica-Bold", 8); canvas.setFillColor(colors.HexColor("#0057B8")); canvas.drawString(20*mm, 285*mm, BRAND.organization)
    canvas.setFont("Helvetica", 7.5); canvas.setFillColor(colors.HexColor("#667085")); canvas.drawString(20*mm, 10*mm, BRAND.product); canvas.drawRightString(width-20*mm, 10*mm, f"Página {doc.page-1}"); canvas.restoreState()


def as_pdf(briefing):
    output = io.BytesIO(); styles = _pdf_styles()
    doc = BaseDocTemplate(output, pagesize=A4, leftMargin=20*mm, rightMargin=20*mm, topMargin=22*mm, bottomMargin=20*mm, title=f"Briefing - {briefing.opportunity.title}", author=BRAND.organization)
    doc.addPageTemplates(PageTemplate(id="3DS", frames=(Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="body"),), onPage=_pdf_page))
    story = [Spacer(1, 24*mm)]
    if LOGO_PATH.exists():
        logo = Image(str(LOGO_PATH), width=36*mm, height=36*mm); logo.hAlign = "CENTER"; story.extend((logo, Spacer(1, 8*mm)))
    story.extend((Paragraph(BRAND.organization, styles["coverbrand"]), Spacer(1, 10*mm), Paragraph("BRIEFING PROFISSIONAL", styles["covertitle"]), Paragraph(clean_text(briefing.opportunity.title), styles["coverproject"]), Spacer(1, 12*mm)))
    rows = [[Paragraph(label.upper(), styles["small"]), Paragraph(value, styles["body"])] for label, value in _metadata(briefing)[2:]]
    if rows:
        table = Table(rows, colWidths=(35*mm, 90*mm), hAlign="CENTER"); table.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("LINEBELOW",(0,0),(-1,-1),.3,colors.HexColor("#D8E4F2")),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)])); story.append(table)
    story.extend((Spacer(1, 18*mm), Paragraph("Dados • Conhecimento • Impacto", styles["coverbrand"]), PageBreak()))
    for section in briefing_sections(briefing):
        content = [Paragraph(section.title, styles["section"])]
        if isinstance(section.value, tuple):
            content.append(ListFlowable([ListItem(Paragraph(item, styles["body"]), leftIndent=5) for item in section.items], bulletType="bullet", leftIndent=16, bulletFontName="Helvetica", bulletFontSize=7, spaceAfter=7))
        elif section.kind in {"summary", "highlight", "alert", "risk"}:
            fill = {"summary":"#EAF2FB","highlight":"#EAF2FB","alert":"#FFF7E8","risk":"#FDEDEE"}[section.kind]
            box = Table([[Paragraph(section.value, styles["body"])]], colWidths=(doc.width-4,)); box.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor(fill)),("BOX",(0,0),(-1,-1),.6,colors.HexColor("#E63946" if section.kind=="risk" else "#0057B8")),("LEFTPADDING",(0,0),(-1,-1),10),("RIGHTPADDING",(0,0),(-1,-1),10),("TOPPADDING",(0,0),(-1,-1),8),("BOTTOMPADDING",(0,0),(-1,-1),8)])); content.append(box)
        else: content.append(Paragraph(section.value, styles["body"]))
        story.append(KeepTogether(content))
    doc.build(story); return output.getvalue()


def _shade_cell(cell, fill):
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr(); tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None: tc_mar = OxmlElement("w:tcMar"); tc_pr.append(tc_mar)
    for margin, value in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = OxmlElement(f"w:{margin}"); node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa"); tc_mar.append(node)


def _field_run(paragraph, instruction):
    run=paragraph.add_run(); begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin"); run._r.append(begin)
    run=paragraph.add_run(); instr=OxmlElement("w:instrText"); instr.set(qn("xml:space"),"preserve"); instr.text=instruction; run._r.append(instr)
    run=paragraph.add_run(); end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end"); run._r.append(end)


def _configure_docx_styles(doc):
    colors_map={"Title":BRAND.graphite,"Subtitle":BRAND.primary,"Heading 1":BRAND.primary,"Heading 2":BRAND.graphite,"Heading 3":BRAND.muted,"Normal":BRAND.graphite}; sizes={"Title":28,"Subtitle":14,"Heading 1":17,"Heading 2":13,"Heading 3":11,"Normal":10.5}
    for name,size in sizes.items():
        style=doc.styles[name]; style.font.name=BRAND.font; style.font.size=Pt(size); style.font.color.rgb=RGBColor.from_string(colors_map[name]); style._element.rPr.rFonts.set(qn("w:ascii"),BRAND.font); style._element.rPr.rFonts.set(qn("w:hAnsi"),BRAND.font)
    doc.styles["Heading 1"].font.bold=True; doc.styles["Heading 1"].paragraph_format.space_before=Pt(15); doc.styles["Heading 1"].paragraph_format.space_after=Pt(7); doc.styles["Heading 1"].paragraph_format.keep_with_next=True
    doc.styles["Normal"].paragraph_format.space_after=Pt(6); doc.styles["Normal"].paragraph_format.line_spacing=1.15


def _docx_header_footer(section):
    header=section.header.paragraphs[0]; header.text=BRAND.organization; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs: run.font.name=BRAND.font; run.font.size=Pt(8); run.font.bold=True; run.font.color.rgb=RGBColor.from_string(BRAND.primary)
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=footer.add_run(f"{BRAND.product}  •  "); run.font.name=BRAND.font; run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string(BRAND.muted); _field_run(footer,"PAGE")


def as_docx(briefing):
    doc=Document(); _configure_docx_styles(doc); cover=doc.sections[0]
    for section in (cover,): section.page_height=Inches(11.7); section.page_width=Inches(8.3); section.top_margin=section.bottom_margin=Inches(.75); section.left_margin=section.right_margin=Inches(.85)
    cover.different_first_page_header_footer=True
    if LOGO_PATH.exists():
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(24); p.add_run().add_picture(str(LOGO_PATH),width=Inches(1.35))
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(14); r=p.add_run(BRAND.organization); r.bold=True; r.font.name=BRAND.font; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(BRAND.primary)
    p=doc.add_paragraph("BRIEFING PROFISSIONAL",style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(34); p.paragraph_format.space_after=Pt(8)
    p=doc.add_paragraph(clean_text(briefing.opportunity.title),style="Subtitle"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    meta=_metadata(briefing)[2:]
    if meta:
        table=doc.add_table(rows=len(meta),cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
        for row,(label,value) in zip(table.rows,meta):
            row.cells[0].width=Inches(1.35); row.cells[1].width=Inches(4.7); row.cells[0].text=label.upper(); row.cells[1].text=value; _shade_cell(row.cells[0],BRAND.pale_blue)
            for cell in row.cells: cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; _set_cell_margins(cell)
            for run in row.cells[0].paragraphs[0].runs: run.font.size=Pt(8); run.font.bold=True; run.font.color.rgb=RGBColor.from_string(BRAND.primary)
    p=doc.add_paragraph("Dados • Conhecimento • Impacto"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(38)
    for run in p.runs: run.font.size=Pt(9); run.font.bold=True; run.font.color.rgb=RGBColor.from_string(BRAND.primary)
    section=doc.add_section(WD_SECTION.NEW_PAGE); section.page_height=Inches(11.7); section.page_width=Inches(8.3); section.top_margin=section.bottom_margin=Inches(.75); section.left_margin=section.right_margin=Inches(.85); _docx_header_footer(section)
    for item in briefing_sections(briefing):
        doc.add_paragraph(item.title,style="Heading 1")
        if isinstance(item.value,tuple):
            for value in item.items: p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(4); p.add_run(value)
        elif item.kind in {"summary","highlight","alert","risk"}:
            table=doc.add_table(rows=1,cols=1); table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.LEFT; cell=table.cell(0,0); cell.text=item.value; _set_cell_margins(cell,130,180,130,180); _shade_cell(cell,BRAND.pale_blue if item.kind in {"summary","highlight"} else ("FFF4DB" if item.kind=="alert" else "FDEDEE"))
        else: doc.add_paragraph(item.value)
    props=doc.core_properties; props.title=f"Briefing - {briefing.opportunity.title}"; props.author=BRAND.organization; props.subject=BRAND.product; props.keywords="briefing, 3DS"
    output=io.BytesIO(); doc.save(output); return output.getvalue()


def _ppt_color(value): return PptRGBColor.from_string(value)


def _ppt_textbox(slide,text,x,y,w,h,size=20,color=None,bold=False,align=PP_ALIGN.LEFT):
    shape=slide.shapes.add_textbox(PptInches(x),PptInches(y),PptInches(w),PptInches(h)); frame=shape.text_frame; frame.clear(); frame.word_wrap=True; frame.margin_left=frame.margin_right=0; paragraph=frame.paragraphs[0]; paragraph.text=text; paragraph.alignment=align; paragraph.font.name=BRAND.font; paragraph.font.size=PptPt(size); paragraph.font.bold=bold; paragraph.font.color.rgb=_ppt_color(color or BRAND.graphite); return shape


def _ppt_footer(slide,number):
    line=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,PptInches(.55),PptInches(7.12),PptInches(12.23),PptInches(.02)); line.fill.solid(); line.fill.fore_color.rgb=_ppt_color("D8E4F2"); line.line.fill.background(); _ppt_textbox(slide,BRAND.organization,.58,7.18,5.8,.22,8,BRAND.muted); _ppt_textbox(slide,str(number),12.1,7.18,.65,.22,8,BRAND.muted,align=PP_ALIGN.RIGHT)


def _ppt_base(prs,title,number):
    slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=_ppt_color("FFFFFF"); accent=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,PptInches(.55),PptInches(.55),PptInches(.1),PptInches(.64)); accent.fill.solid(); accent.fill.fore_color.rgb=_ppt_color(BRAND.primary); accent.line.fill.background(); _ppt_textbox(slide,title,.82,.5,11.3,.75,35,BRAND.graphite,True); _ppt_footer(slide,number); return slide


def _ppt_bullets(slide,items):
    shape=slide.shapes.add_textbox(PptInches(.85),PptInches(1.55),PptInches(11.6),PptInches(5.2)); frame=shape.text_frame; frame.clear(); frame.word_wrap=True; frame.margin_left=frame.margin_right=PptInches(.08); frame.vertical_anchor=MSO_ANCHOR.TOP
    for index,item in enumerate(items):
        p=frame.paragraphs[0] if index==0 else frame.add_paragraph(); p.text=f"• {item}"; p.font.name=BRAND.font; p.font.size=PptPt(20); p.font.color.rgb=_ppt_color(BRAND.graphite); p.space_after=PptPt(14)


def as_pptx(briefing):
    prs=Presentation(); prs.slide_width=PptInches(13.333); prs.slide_height=PptInches(7.5); cover=prs.slides.add_slide(prs.slide_layouts[6]); cover.background.fill.solid(); cover.background.fill.fore_color.rgb=_ppt_color(BRAND.graphite)
    bar=cover.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,PptInches(.16),prs.slide_height); bar.fill.solid(); bar.fill.fore_color.rgb=_ppt_color(BRAND.primary); bar.line.fill.background()
    if LOGO_PATH.exists(): cover.shapes.add_picture(str(LOGO_PATH),PptInches(10.65),PptInches(.55),width=PptInches(1.8),height=PptInches(1.8))
    _ppt_textbox(cover,BRAND.organization,.8,.65,7.8,.4,14,"FFFFFF",True); _ppt_textbox(cover,"Briefing profissional",.8,2.0,10.6,.85,50,"FFFFFF",True); _ppt_textbox(cover,clean_text(briefing.opportunity.title),.82,3.0,10.5,.85,26,"65B5FF")
    meta="  •  ".join(filter(None,(clean_text(briefing.opportunity.client_name),_document_date(briefing))))
    if meta: _ppt_textbox(cover,meta,.82,5.75,10.5,.4,15,"D0D5DD")
    sections={item.key:item for item in briefing_sections(briefing)}; number=2
    groups=(("O briefing alinha o desafio e o resultado esperado",("problem","objective","target_audience")),("A solução é definida por entregas e requisitos verificáveis",("deliverables","functional_requirements","non_functional_requirements")),("A execução depende de uma base técnica explicitamente alinhada",("integrations","data_requirements","infrastructure_requirements")),("Premissas e riscos precisam ser resolvidos antes da entrega",("constraints","dependencies","assumptions","scope_risks","ambiguities")),("O próximo passo é validar critérios e fechar as lacunas",("acceptance_criteria","client_questions")))
    for title,keys in groups:
        if not any(key in sections for key in keys): continue
        slide=_ppt_base(prs,title,number); number+=1; items=[]
        for key in keys:
            if key in sections: items.extend(f"{sections[key].title}: {value}" for value in sections[key].items)
        _ppt_bullets(slide,items[:8])
    props=prs.core_properties; props.title=f"Briefing - {briefing.opportunity.title}"; props.author=BRAND.organization; props.subject=BRAND.product; props.keywords="briefing, 3DS"
    output=io.BytesIO(); prs.save(output); return output.getvalue()


def render_export(briefing, export_format):
    if export_format=="json": return json.dumps(public_payload(briefing),ensure_ascii=False,indent=2).encode("utf-8")
    if export_format=="md": return as_markdown(briefing)
    if export_format=="pdf": return as_pdf(briefing)
    if export_format=="docx": return as_docx(briefing)
    if export_format=="pptx": return as_pptx(briefing)
    raise ValueError("Formato de exportação inválido.")


def _proposal_metadata(proposal):
    opportunity = proposal.opportunity
    versions = {"SHORT": "Curta", "CONSULTATIVE": "Consultiva", "TECHNICAL": "Técnica"}
    result = [("Documento", "Proposta Comercial")]
    if clean_text(opportunity.client_name):
        result.append(("Cliente", clean_text(opportunity.client_name)))
    result.extend((
        ("Projeto", clean_text(opportunity.title)),
        ("Versão", versions.get(proposal.version, proposal.version)),
    ))
    if proposal.suggested_price is not None:
        amount = f"{proposal.suggested_price:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
        result.append(("Investimento", f"{clean_text(proposal.currency)} {amount}"))
    result.append(("Data", _document_date(proposal)))
    return result


def as_proposal_pdf(proposal):
    output = io.BytesIO()
    styles = _pdf_styles()
    doc = BaseDocTemplate(
        output, pagesize=A4, leftMargin=20*mm, rightMargin=20*mm,
        topMargin=22*mm, bottomMargin=20*mm,
        title=f"Proposta Comercial - {proposal.opportunity.title}", author=BRAND.organization,
    )
    doc.addPageTemplates(PageTemplate(id="3DS", frames=(Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="body"),), onPage=_pdf_page))
    story = [Spacer(1, 24*mm)]
    if LOGO_PATH.exists():
        logo = Image(str(LOGO_PATH), width=36*mm, height=36*mm)
        logo.hAlign = "CENTER"
        story.extend((logo, Spacer(1, 8*mm)))
    story.extend((
        Paragraph(BRAND.organization, styles["coverbrand"]), Spacer(1, 10*mm),
        Paragraph("PROPOSTA COMERCIAL", styles["covertitle"]),
        Paragraph(xml_escape(clean_text(proposal.opportunity.title)), styles["coverproject"]), Spacer(1, 12*mm),
    ))
    rows = [[Paragraph(xml_escape(label.upper()), styles["small"]), Paragraph(xml_escape(value), styles["body"])] for label, value in _proposal_metadata(proposal)[1:]]
    if rows:
        table = Table(rows, colWidths=(35*mm, 90*mm), hAlign="CENTER")
        table.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("LINEBELOW",(0,0),(-1,-1),.3,colors.HexColor("#D8E4F2")),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]))
        story.append(table)
    story.extend((Spacer(1, 18*mm), Paragraph("Dados • Conhecimento • Impacto", styles["coverbrand"]), PageBreak()))
    for section in proposal_sections(proposal):
        content = [Paragraph(section.title, styles["section"])]
        if isinstance(section.value, tuple):
            content.append(ListFlowable([ListItem(Paragraph(xml_escape(item), styles["body"]), leftIndent=5) for item in section.items], bulletType="bullet", leftIndent=16, bulletFontName="Helvetica", bulletFontSize=7, spaceAfter=7))
        elif section.kind in {"summary", "highlight"}:
            box = Table([[Paragraph(xml_escape(section.value), styles["body"])]], colWidths=(doc.width-4,))
            box.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#EAF2FB")),("BOX",(0,0),(-1,-1),.6,colors.HexColor("#0057B8")),("LEFTPADDING",(0,0),(-1,-1),10),("RIGHTPADDING",(0,0),(-1,-1),10),("TOPPADDING",(0,0),(-1,-1),8),("BOTTOMPADDING",(0,0),(-1,-1),8)]))
            content.append(box)
        else:
            content.append(Paragraph(xml_escape(section.value), styles["body"]))
        story.append(KeepTogether(content))
    doc.build(story)
    return output.getvalue()


def as_proposal_docx(proposal):
    doc = Document()
    _configure_docx_styles(doc)
    cover = doc.sections[0]
    cover.page_height = Inches(11.7); cover.page_width = Inches(8.3)
    cover.top_margin = cover.bottom_margin = Inches(.75); cover.left_margin = cover.right_margin = Inches(.85)
    cover.different_first_page_header_footer = True
    if LOGO_PATH.exists():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(24); p.add_run().add_picture(str(LOGO_PATH), width=Inches(1.35))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(14)
    run = p.add_run(BRAND.organization); run.bold = True; run.font.name = BRAND.font; run.font.size = Pt(11); run.font.color.rgb = RGBColor.from_string(BRAND.primary)
    p = doc.add_paragraph("PROPOSTA COMERCIAL", style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(34)
    p = doc.add_paragraph(clean_text(proposal.opportunity.title), style="Subtitle"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = _proposal_metadata(proposal)[1:]
    if meta:
        table = doc.add_table(rows=len(meta), cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
        for row, (label, value) in zip(table.rows, meta):
            row.cells[0].width = Inches(1.35); row.cells[1].width = Inches(4.7)
            row.cells[0].text = label.upper(); row.cells[1].text = value; _shade_cell(row.cells[0], BRAND.pale_blue)
            for cell in row.cells: cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; _set_cell_margins(cell)
            for text_run in row.cells[0].paragraphs[0].runs: text_run.font.size = Pt(8); text_run.font.bold = True; text_run.font.color.rgb = RGBColor.from_string(BRAND.primary)
    p = doc.add_paragraph("Dados • Conhecimento • Impacto"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(38)
    for text_run in p.runs: text_run.font.size = Pt(9); text_run.font.bold = True; text_run.font.color.rgb = RGBColor.from_string(BRAND.primary)
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_height = Inches(11.7); section.page_width = Inches(8.3)
    section.top_margin = section.bottom_margin = Inches(.75); section.left_margin = section.right_margin = Inches(.85)
    _docx_header_footer(section)
    for item in proposal_sections(proposal):
        doc.add_paragraph(item.title, style="Heading 1")
        if isinstance(item.value, tuple):
            for value in item.items:
                doc.add_paragraph(value, style="List Bullet")
        elif item.kind in {"summary", "highlight"}:
            table = doc.add_table(rows=1, cols=1); cell = table.cell(0, 0); cell.text = item.value
            _set_cell_margins(cell, 130, 180, 130, 180); _shade_cell(cell, BRAND.pale_blue)
        else:
            doc.add_paragraph(item.value)
    props = doc.core_properties
    props.title = f"Proposta Comercial - {proposal.opportunity.title}"
    props.author = BRAND.organization; props.subject = BRAND.product; props.keywords = "proposta comercial, 3DS"
    output = io.BytesIO(); doc.save(output)
    return output.getvalue()


def render_proposal_export(proposal, export_format):
    if export_format == "pdf":
        return as_proposal_pdf(proposal)
    if export_format == "docx":
        return as_proposal_docx(proposal)
    raise ValueError("Formato de exportação de proposta inválido.")
