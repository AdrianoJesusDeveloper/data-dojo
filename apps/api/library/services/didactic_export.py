from html import escape
from io import BytesIO
import re

from docx import Document
from docx.shared import Pt

from ..models import DidacticLesson


UNSOURCED_WARNING = (
    "Aula gerada por IA sem fonte aprovada. Este conteúdo permanece sujeito a "
    "revisão humana e não deve ser publicado automaticamente."
)


def build_export_snapshot(lesson: DidacticLesson) -> dict:
    unit = lesson.learning_target
    module = unit.module
    formation = module.formation
    return {
        "lesson_id": lesson.pk, "title": lesson.title,
        "formation": formation.title, "formation_id": formation.pk,
        "module": module.title, "module_id": module.pk,
        "unit": unit.title, "unit_id": unit.pk,
        "status": lesson.get_status_display(), "status_code": lesson.status,
        "audience": lesson.get_audience_display(), "audience_code": lesson.audience,
        "source_mode": lesson.get_source_mode_display(), "source_mode_code": lesson.source_mode,
        "ai_provider": lesson.ai_provider, "ai_model": lesson.ai_model,
        "generated_at": lesson.generated_at, "updated_at": lesson.updated_at,
        "warning": UNSOURCED_WARNING if lesson.source_mode == DidacticLesson.SourceMode.AI_GENERATED_UNSOURCED else "",
        "sections": [{"type": item.get_section_type_display(), "title": item.title, "content": item.content} for item in lesson.sections.all()],
        "sources": [{"title": item.title, "reference": item.reference, "location": item.location, "url": item.url} for item in lesson.sources.all()],
    }


def export_filename(snapshot: dict, extension: str) -> str:
    stem = re.sub(r"[^a-zA-Z0-9_-]+", "-", snapshot["title"]).strip("-").lower() or "aula-didatica"
    return f"{stem}-{snapshot['lesson_id']}.{extension}"


def render_lesson_docx(snapshot: dict) -> bytes:
    document = Document()
    document.core_properties.title = snapshot["title"]
    document.core_properties.subject = "Conteúdo Didático — Data Driven Dojô"
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(snapshot["title"], level=0)
    for label, value in _metadata(snapshot):
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(value)
    if snapshot["warning"]:
        document.add_paragraph().add_run(snapshot["warning"]).bold = True
    for section in snapshot["sections"]:
        document.add_heading(section["title"], level=1)
        document.add_paragraph(section["type"], style="Subtitle")
        for block in section["content"].splitlines() or [""]:
            document.add_paragraph(block)
    if snapshot["sources"]:
        document.add_heading("Fontes e referências", level=1)
        for source in snapshot["sources"]:
            details = " — ".join(filter(None, (source["reference"], source["location"], source["url"])))
            document.add_paragraph(f"{source['title']}{': ' + details if details else ''}", style="List Bullet")
    document.add_paragraph(_traceability(snapshot))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_lesson_html(snapshot: dict) -> bytes:
    metadata = "".join(f"<dt>{escape(label)}</dt><dd>{escape(value)}</dd>" for label, value in _metadata(snapshot))
    warning = f'<aside class="warning">{escape(snapshot["warning"])}</aside>' if snapshot["warning"] else ""
    sections = "".join(
        f'<section><p class="section-type">{escape(item["type"])}</p><h2>{escape(item["title"])}</h2>'
        f'<div>{escape(item["content"]).replace(chr(10), "<br>")}</div></section>' for item in snapshot["sections"]
    )
    sources = ""
    if snapshot["sources"]:
        items = "".join(f"<li><strong>{escape(item['title'])}</strong> {escape(' — '.join(filter(None, (item['reference'], item['location'], item['url']))))}</li>" for item in snapshot["sources"])
        sources = f"<section><h2>Fontes e referências</h2><ul>{items}</ul></section>"
    html = f'''<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{escape(snapshot['title'])}</title><style>
body{{max-width:900px;margin:40px auto;padding:0 24px;color:#171717;font:16px/1.6 Arial,sans-serif}}h1{{border-bottom:2px solid #171717;padding-bottom:12px}}section{{margin:32px 0}}dl{{display:grid;grid-template-columns:max-content 1fr;gap:4px 16px}}dt{{font-weight:700}}dd{{margin:0}}.section-type{{font-size:12px;text-transform:uppercase;color:#555}}.warning{{padding:14px;border:1px solid #a16207;background:#fef3c7;color:#713f12;font-weight:700}}footer{{border-top:1px solid #bbb;margin-top:40px;padding-top:12px;color:#555;font-size:13px}}
</style></head><body><main><h1>{escape(snapshot['title'])}</h1><dl>{metadata}</dl>{warning}{sections}{sources}</main><footer>{escape(_traceability(snapshot))}</footer></body></html>'''
    return html.encode("utf-8")


def _metadata(snapshot: dict) -> list[tuple[str, str]]:
    values = [
        ("Formação", snapshot["formation"]), ("Módulo", snapshot["module"]), ("Unidade", snapshot["unit"]),
        ("Status", f"{snapshot['status']} ({snapshot['status_code']})"),
        ("Audiência", f"{snapshot['audience']} ({snapshot['audience_code']})"),
        ("Modo de fonte", f"{snapshot['source_mode']} ({snapshot['source_mode_code']})"),
    ]
    if snapshot["ai_provider"]:
        values.append(("Proveniência IA", snapshot["ai_provider"] + (f" / {snapshot['ai_model']}" if snapshot["ai_model"] else "")))
    if snapshot["generated_at"]:
        values.append(("Gerada em", snapshot["generated_at"].isoformat()))
    return values


def _traceability(snapshot: dict) -> str:
    return (f"Rastreabilidade Data Driven Dojô — aula #{snapshot['lesson_id']}; formação #{snapshot['formation_id']}; "
            f"módulo #{snapshot['module_id']}; unidade #{snapshot['unit_id']}; versão atualizada em {snapshot['updated_at'].isoformat()}.")
