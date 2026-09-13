import io
from dataclasses import dataclass
from datetime import date
from xml.sax.saxutils import escape as xml_escape

from django.utils.text import slugify
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.util import Inches as PptInches, Pt as PptPt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate, Frame, Image, KeepTogether, ListFlowable, ListItem, PageBreak,
    PageTemplate, Paragraph, Spacer, Table, TableStyle,
)

from professional.document_branding import BRAND as PROFESSIONAL_BRAND, LOGO_PATH
from library.editorial_contracts import normalize_project_type


@dataclass(frozen=True)
class ContentStudioBrand:
    organization: str = PROFESSIONAL_BRAND.organization
    product: str = "Content Studio"
    primary: str = PROFESSIONAL_BRAND.primary
    graphite: str = PROFESSIONAL_BRAND.graphite
    accent: str = PROFESSIONAL_BRAND.accent
    risk: str = PROFESSIONAL_BRAND.risk
    muted: str = PROFESSIONAL_BRAND.muted
    pale_blue: str = PROFESSIONAL_BRAND.pale_blue
    pale_gray: str = PROFESSIONAL_BRAND.pale_gray
    font: str = PROFESSIONAL_BRAND.font


BRAND = ContentStudioBrand()

MIMES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def clean_text(value):
    # Non-breaking/soft hyphens commonly arrive from copied editorial text.
    text = str(value or "").translate(str.maketrans({"\u2011": "-", "\u2010": "-", "\u00ad": "", "\u200b": ""}))
    return " ".join(text.split()).strip()


def _humanize_key(value):
    return clean_text(value).replace("_", " ").strip().capitalize()


def _structured_text(value):
    """Turn structured editorial values into readable document text."""
    if value is None:
        return ""
    if isinstance(value, dict):
        language = clean_text(value.get("language"))
        snippet = value.get("snippet")
        if snippet is not None:
            code = str(snippet).strip()
            return f"{language}:\n{code}" if language else code
        parts = []
        for key, item in value.items():
            rendered = _structured_text(item)
            if rendered:
                parts.append(f"{_humanize_key(key)}: {rendered}")
        return " | ".join(parts)
    if isinstance(value, (list, tuple)):
        return "; ".join(filter(None, (_structured_text(item) for item in value)))
    return clean_text(value)


def clean_items(value):
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        items = value
    else:
        items = [value]
    result = []
    for item in items:
        if isinstance(item, dict):
            text = clean_text(
                item.get("title")
                or item.get("name")
                or item.get("theme")
                or item.get("objective")
                or item
            )
        else:
            text = clean_text(item)
        if text:
            result.append(text)
    return result


def _document_date(project):
    value = getattr(project, "updated_at", None)
    return (value.date() if value else date.today()).strftime("%d/%m/%Y")


def export_filename(project, extension):
    base = slugify(project.title) or f"projeto-{project.pk}"
    return f"content-studio-{base}.{extension}"


def _mode_label(project):
    semantic_project_type = normalize_project_type(project.project_type)
    return "Conteúdo Editorial" if semantic_project_type == "content" else "Formação Premium"


def _research_label(project):
    labels = {
        "ACERVO_ONLY": "Somente acervo",
        "WEB_ONLY": "Somente web",
        "HYBRID": "Acervo + web",
    }
    return labels.get(project.research_policy, project.research_policy)


def _source_rows(project, plan):
    rows = []
    seen = set()

    try:
        for citation in project.citations.all():
            title = clean_text(citation.book_title) or "Fonte do acervo"
            page = f"p. {citation.page_number}" if citation.page_number else ""
            excerpt = clean_text(citation.excerpt)
            key = ("ACERVO", title, page, excerpt)
            if key in seen:
                continue
            seen.add(key)
            rows.append({
                "kind": "ACERVO",
                "title": title,
                "reference": page,
                "excerpt": excerpt,
            })
    except Exception:
        pass

    try:
        context = project.research_context
        for evidence in context.evidence.all():
            kind = clean_text(evidence.source_kind) or "FONTE"
            title = clean_text(evidence.title) or clean_text(evidence.domain) or "Evidência"
            reference = clean_text(evidence.url) or clean_text(evidence.domain)
            excerpt = clean_text(evidence.excerpt)
            key = (kind, title, reference, excerpt)
            if key in seen:
                continue
            seen.add(key)
            rows.append({
                "kind": kind,
                "title": title,
                "reference": reference,
                "excerpt": excerpt,
            })
    except Exception:
        pass

    architecture = plan.proposed_architecture or {}
    for source in clean_items(architecture.get("sources")):
        key = ("PLANO", source, "", "")
        if key in seen:
            continue
        seen.add(key)
        rows.append({
            "kind": "PLANO",
            "title": source,
            "reference": "",
            "excerpt": "",
        })

    if not rows:
        rows.append({
            "kind": "SEM FONTE",
            "title": "Sem fonte verificada — validar posteriormente",
            "reference": "",
            "excerpt": "",
        })
    return rows


def _metadata(project, plan):
    return [
        ("Documento", "Plano editorial"),
        ("Projeto", clean_text(project.title)),
        ("Formato", _mode_label(project)),
        ("Versão", str(plan.version)),
        ("Estado", clean_text(plan.get_status_display() if hasattr(plan, "get_status_display") else plan.status)),
        ("Pesquisa", _research_label(project)),
        ("Data", _document_date(project)),
    ]


def _authorship_lines(challenge):
    if not isinstance(challenge, dict):
        return []
    labels = (
        ("independent_explanation", "Explicação independente"),
        ("practical_challenge", "Desafio prático"),
        ("portfolio_artifact", "Artefato de portfólio"),
        ("reflection_question", "Reflexão"),
        ("comprehension_criteria", "Critério de compreensão"),
        ("responsible_ai_use", "Uso responsável de IA"),
        ("must_not_delegate_to_ai", "Não delegar à IA"),
        ("expected_result", "Resultado esperado"),
        ("private_submission_option", "Alternativa privada"),
    )
    result = []
    for key, label in labels:
        text = clean_text(challenge.get(key))
        if text:
            result.append(f"{label}: {text}")
    return result


def _youtube_sections(project, plan):
    a = plan.proposed_architecture or {}
    overview = []
    for key, label in (
        ("objective", "Objetivo"),
        ("target_audience", "Público-alvo"),
        ("level", "Nível"),
        ("playlist_description", "Descrição da série"),
        ("estimated_total_duration", "Duração estimada"),
    ):
        text = clean_text(a.get(key))
        if text:
            overview.append(f"{label}: {text}")

    for key, label in (
        ("prerequisites", "Pré-requisitos"),
        ("competencies", "Competências"),
        ("tools", "Ferramentas"),
    ):
        items = clean_items(a.get(key))
        if items:
            overview.append(f"{label}: " + "; ".join(items))

    sections = [("Visão editorial", overview)] if overview else []
    for index, video in enumerate(a.get("videos", []) or [], start=1):
        if not isinstance(video, dict):
            continue
        order = video.get("order", index)
        title = clean_text(video.get("title")) or f"Vídeo {order}"
        lines = []
        for key, label in (
            ("theme", "Tema"),
            ("objective", "Objetivo"),
            ("practical_demo", "Demonstração prática"),
            ("practice", "Prática"),
            ("code", "Código"),
            ("exercise", "Exercício"),
            ("ai_integration", "Integração com IA"),
            ("human_reasoning", "Raciocínio humano"),
            ("validation", "Validação"),
            ("reflection", "Reflexão"),
            ("without_ai_challenge", "Desafio sem IA"),
        ):
            value = video.get(key)
            rendered = _structured_text(value)
            if rendered:
                lines.append(f"{label}: {rendered}")
        for key, label in (("concepts", "Conceitos"), ("tools", "Ferramentas"), ("rag_sources", "Fontes do vídeo")):
            items = clean_items(video.get(key))
            if items:
                lines.append(f"{label}: " + "; ".join(items))
        authorship = _authorship_lines(video.get("authorship_challenge"))
        if authorship:
            lines.append("Desafio de Autoria")
            lines.extend(authorship)
        sections.append((f"Vídeo {order} — {title}", lines))
    return sections


def _premium_sections(project, plan):
    a = plan.proposed_architecture or {}
    overview = []
    for key, label in (
        ("objective", "Objetivo"),
        ("target_audience", "Público-alvo"),
        ("level", "Nível"),
        ("estimated_total_duration", "Duração estimada"),
    ):
        text = clean_text(a.get(key))
        if text:
            overview.append(f"{label}: {text}")

    for key, label in (
        ("prerequisites", "Pré-requisitos"),
        ("competencies", "Competências"),
        ("tools", "Ferramentas"),
    ):
        items = clean_items(a.get(key))
        if items:
            overview.append(f"{label}: " + "; ".join(items))

    sections = [("Visão editorial", overview)] if overview else []
    for m_index, module in enumerate(a.get("modules", []) or [], start=1):
        if not isinstance(module, dict):
            continue
        m_order = module.get("order", m_index)
        m_title = clean_text(module.get("title")) or f"Módulo {m_order}"
        m_lines = []
        for key, label in (("objective", "Objetivo"), ("description", "Descrição")):
            text = clean_text(module.get(key))
            if text:
                m_lines.append(f"{label}: {text}")
        for key, label in (("competencies", "Competências"), ("tools", "Ferramentas")):
            items = clean_items(module.get(key))
            if items:
                m_lines.append(f"{label}: " + "; ".join(items))
        sections.append((f"Módulo {m_order} — {m_title}", m_lines))

        for l_index, lesson in enumerate(module.get("lessons", []) or [], start=1):
            if not isinstance(lesson, dict):
                continue
            l_order = lesson.get("order", l_index)
            l_title = clean_text(lesson.get("title")) or f"Aula {l_order}"
            lines = []
            for key, label in (
                ("theme", "Tema"),
                ("objective", "Objetivo"),
                ("practical_demo", "Demonstração prática"),
                ("practice", "Prática"),
                ("code", "Código"),
                ("exercise", "Exercício"),
                ("ai_integration", "Integração com IA"),
                ("human_reasoning", "Raciocínio humano"),
                ("validation", "Validação"),
                ("reflection", "Reflexão"),
                ("without_ai_challenge", "Desafio sem IA"),
            ):
                value = lesson.get(key)
                rendered = _structured_text(value)
                if rendered:
                    lines.append(f"{label}: {rendered}")
            for key, label in (("concepts", "Conceitos"), ("tools", "Ferramentas"), ("rag_sources", "Fontes da aula")):
                items = clean_items(lesson.get(key))
                if items:
                    lines.append(f"{label}: " + "; ".join(items))
            authorship = _authorship_lines(lesson.get("authorship_challenge"))
            if authorship:
                lines.append("Desafio de Autoria")
                lines.extend(authorship)
            sections.append((f"Módulo {m_order} · Aula {l_order} — {l_title}", lines))

    for key, title in (
        ("final_project", "Projeto final"),
        ("assessment", "Avaliação"),
        ("certification", "Certificação"),
    ):
        value = a.get(key)
        if isinstance(value, dict):
            lines = [f"{clean_text(k).replace('_', ' ').title()}: {clean_text(v)}" for k, v in value.items() if clean_text(v)]
        else:
            lines = clean_items(value)
        if lines:
            sections.append((title, lines))
    return sections


def _sections(project, plan):
    semantic_project_type = normalize_project_type(project.project_type)
    return _youtube_sections(project, plan) if semantic_project_type == "content" else _premium_sections(project, plan)


def _pdf_styles():
    sample = getSampleStyleSheet()
    return {
        "body": ParagraphStyle("CS Body", parent=sample["BodyText"], fontName="Helvetica", fontSize=10.2, leading=14, textColor=colors.HexColor("#1C1C1C"), spaceAfter=5),
        "section": ParagraphStyle("CS Section", parent=sample["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, textColor=colors.HexColor("#0057B8"), spaceBefore=10, spaceAfter=7, keepWithNext=True),
        "small": ParagraphStyle("CS Small", parent=sample["BodyText"], fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor("#667085")),
        "coverbrand": ParagraphStyle("CS Cover Brand", parent=sample["BodyText"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=colors.HexColor("#0057B8"), alignment=1),
        "covertitle": ParagraphStyle("CS Cover Title", parent=sample["Title"], fontName="Helvetica-Bold", fontSize=28, leading=33, textColor=colors.HexColor("#1C1C1C"), alignment=1, spaceAfter=10),
        "coverproject": ParagraphStyle("CS Cover Project", parent=sample["Heading2"], fontName="Helvetica", fontSize=16, leading=21, textColor=colors.HexColor("#0057B8"), alignment=1),
    }


def _pdf_page(canvas, doc):
    if doc.page == 1:
        return
    canvas.saveState()
    width, _ = A4
    canvas.setStrokeColor(colors.HexColor("#D8E4F2"))
    canvas.line(20 * mm, 281 * mm, width - 20 * mm, 281 * mm)
    canvas.line(20 * mm, 15 * mm, width - 20 * mm, 15 * mm)
    canvas.setFont("Helvetica-Bold", 8)
    canvas.setFillColor(colors.HexColor("#0057B8"))
    canvas.drawString(20 * mm, 285 * mm, BRAND.organization)
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#667085"))
    canvas.drawString(20 * mm, 10 * mm, BRAND.product)
    canvas.drawRightString(width - 20 * mm, 10 * mm, f"Página {doc.page - 1}")
    canvas.restoreState()


def render_plan_pdf(project, plan):
    output = io.BytesIO()
    styles = _pdf_styles()
    doc = BaseDocTemplate(
        output, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=22 * mm, bottomMargin=20 * mm,
        title=f"Plano editorial - {project.title}", author=BRAND.organization,
    )
    doc.addPageTemplates(PageTemplate(id="3DS-CS", frames=(Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="body"),), onPage=_pdf_page))
    story = [Spacer(1, 24 * mm)]
    if LOGO_PATH.exists():
        logo = Image(str(LOGO_PATH), width=36 * mm, height=36 * mm)
        logo.hAlign = "CENTER"
        story.extend((logo, Spacer(1, 8 * mm)))
    story.extend((
        Paragraph(BRAND.organization, styles["coverbrand"]),
        Spacer(1, 10 * mm),
        Paragraph("CONTENT STUDIO", styles["covertitle"]),
        Paragraph(xml_escape(clean_text(project.title)), styles["coverproject"]),
        Spacer(1, 12 * mm),
    ))
    rows = [[Paragraph(xml_escape(label.upper()), styles["small"]), Paragraph(xml_escape(value), styles["body"])] for label, value in _metadata(project, plan)[2:]]
    table = Table(rows, colWidths=(35 * mm, 90 * mm), hAlign="CENTER")
    table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 0), (-1, -1), .3, colors.HexColor("#D8E4F2")),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend((table, Spacer(1, 18 * mm), Paragraph("Aprender • Construir • Publicar • Ensinar", styles["coverbrand"]), PageBreak()))

    for title, lines in _sections(project, plan):
        content = [Paragraph(xml_escape(title), styles["section"])]
        if lines:
            content.append(ListFlowable(
                [ListItem(Paragraph(xml_escape(line), styles["body"]), leftIndent=5) for line in lines],
                bulletType="bullet", leftIndent=16, bulletFontName="Helvetica", bulletFontSize=7, spaceAfter=7,
            ))
        story.append(KeepTogether(content))

    story.append(PageBreak())
    story.append(Paragraph("Fontes e proveniência", styles["section"]))
    for source in _source_rows(project, plan):
        line = f"[{source['kind']}] {source['title']}"
        if source["reference"]:
            line += f" — {source['reference']}"
        story.append(Paragraph(xml_escape(line), styles["body"]))
        if source["excerpt"]:
            story.append(Paragraph(xml_escape(source["excerpt"]), styles["small"]))

    doc.build(story)
    return output.getvalue()


def _shade_cell(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{margin}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def _field_run(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _configure_docx_styles(doc):
    colors_map = {"Title": BRAND.graphite, "Subtitle": BRAND.primary, "Heading 1": BRAND.primary, "Heading 2": BRAND.graphite, "Heading 3": BRAND.muted, "Normal": BRAND.graphite}
    sizes = {"Title": 28, "Subtitle": 14, "Heading 1": 17, "Heading 2": 13, "Heading 3": 11, "Normal": 10.5}
    for name, size in sizes.items():
        style = doc.styles[name]
        style.font.name = BRAND.font
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(colors_map[name])
        style._element.rPr.rFonts.set(qn("w:ascii"), BRAND.font)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BRAND.font)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(15)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15


def _docx_header_footer(section):
    header = section.header.paragraphs[0]
    header.text = BRAND.organization
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = BRAND.font
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(BRAND.primary)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"{BRAND.product}  •  ")
    run.font.name = BRAND.font
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(BRAND.muted)
    _field_run(footer, "PAGE")


def render_plan_docx(project, plan):
    doc = Document()
    _configure_docx_styles(doc)
    cover = doc.sections[0]
    cover.page_height = Inches(11.7)
    cover.page_width = Inches(8.3)
    cover.top_margin = cover.bottom_margin = Inches(.75)
    cover.left_margin = cover.right_margin = Inches(.85)
    cover.different_first_page_header_footer = True

    if LOGO_PATH.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(24)
        p.add_run().add_picture(str(LOGO_PATH), width=Inches(1.35))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    r = p.add_run(BRAND.organization)
    r.bold = True
    r.font.name = BRAND.font
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(BRAND.primary)

    p = doc.add_paragraph("CONTENT STUDIO", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph(clean_text(project.title), style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = _metadata(project, plan)[2:]
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row, (label, value) in zip(table.rows, meta):
        row.cells[0].width = Inches(1.35)
        row.cells[1].width = Inches(4.7)
        row.cells[0].text = label.upper()
        row.cells[1].text = value
        _shade_cell(row.cells[0], BRAND.pale_blue)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
        for run in row.cells[0].paragraphs[0].runs:
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(BRAND.primary)

    p = doc.add_paragraph("Aprender • Construir • Publicar • Ensinar")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(38)
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(BRAND.primary)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_height = Inches(11.7)
    section.page_width = Inches(8.3)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    _docx_header_footer(section)

    for title, lines in _sections(project, plan):
        doc.add_paragraph(title, style="Heading 1")
        for line in lines:
            doc.add_paragraph(line, style="List Bullet")

    doc.add_paragraph("Fontes e proveniência", style="Heading 1")
    for source in _source_rows(project, plan):
        p = doc.add_paragraph(style="List Bullet")
        text = f"[{source['kind']}] {source['title']}"
        if source["reference"]:
            text += f" — {source['reference']}"
        p.add_run(text)
        if source["excerpt"]:
            q = doc.add_paragraph(source["excerpt"])
            q.paragraph_format.left_indent = Inches(.25)

    props = doc.core_properties
    props.title = f"Plano editorial - {project.title}"
    props.author = BRAND.organization
    props.subject = BRAND.product
    props.keywords = "content studio, plano editorial, 3DS"

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _ppt_color(value):
    return PptRGBColor.from_string(value)


def _ppt_textbox(slide, text, x, y, w, h, size=20, color=None, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    p = frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = BRAND.font
    p.font.size = PptPt(size)
    p.font.bold = bold
    p.font.color.rgb = _ppt_color(color or BRAND.graphite)
    return shape


def _ppt_footer(slide, number):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(.55), PptInches(7.12), PptInches(12.23), PptInches(.02))
    line.fill.solid()
    line.fill.fore_color.rgb = _ppt_color("D8E4F2")
    line.line.fill.background()
    _ppt_textbox(slide, BRAND.organization, .58, 7.18, 5.8, .22, 8, BRAND.muted)
    _ppt_textbox(slide, str(number), 12.1, 7.18, .65, .22, 8, BRAND.muted, align=PP_ALIGN.RIGHT)


def _ppt_base(prs, title, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _ppt_color("FFFFFF")
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(.55), PptInches(.55), PptInches(.1), PptInches(.64))
    accent.fill.solid()
    accent.fill.fore_color.rgb = _ppt_color(BRAND.primary)
    accent.line.fill.background()
    _ppt_textbox(slide, title, .82, .5, 11.3, .75, 31, BRAND.graphite, True)
    _ppt_footer(slide, number)
    return slide


def _ppt_bullets(slide, items, max_items=7):
    items = [clean_text(item) for item in items if clean_text(item)][:max_items]
    shape = slide.shapes.add_textbox(PptInches(.85), PptInches(1.5), PptInches(11.6), PptInches(5.3))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = PptInches(.08)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = f"• {item}"
        p.font.name = BRAND.font
        p.font.size = PptPt(18 if len(items) > 5 else 20)
        p.font.color.rgb = _ppt_color(BRAND.graphite)
        p.space_after = PptPt(10)


def _split_for_slides(title, lines):
    if not lines:
        return [(title, ["Conteúdo editorial disponível no plano."])]
    chunks = []
    size = 6
    for i in range(0, len(lines), size):
        chunk_title = title if i == 0 else f"{title} — continuação"
        chunks.append((chunk_title, lines[i:i + size]))
    return chunks


def _split_labeled_lines(lines):
    groups = {
        "context": [],
        "practice": [],
        "human_ai": [],
        "authorship": [],
        "sources": [],
    }
    authorship = False
    for raw in lines:
        line = clean_text(raw)
        if not line:
            continue
        if line == "Desafio de Autoria":
            authorship = True
            continue
        if authorship:
            groups["authorship"].append(line)
            continue
        prefix = line.split(":", 1)[0].lower()
        if prefix in {"fontes do vídeo", "fontes da aula"}:
            groups["sources"].append(line)
        elif prefix in {"demonstração prática", "prática", "código", "exercício", "desafio sem ia"}:
            groups["practice"].append(line)
        elif prefix in {"integração com ia", "raciocínio humano", "validação", "reflexão"}:
            groups["human_ai"].append(line)
        else:
            groups["context"].append(line)
    return groups


def _semantic_slide_sections(title, lines):
    """Transform a dense video/lesson block into editorially meaningful slides."""
    groups = _split_labeled_lines(lines)
    result = []
    if groups["context"]:
        result.append((title, groups["context"]))
    if groups["practice"]:
        result.append((f"{title} — Demonstração e prática", groups["practice"]))
    if groups["human_ai"]:
        result.append((f"{title} — IA, raciocínio humano e validação", groups["human_ai"]))
    if groups["authorship"]:
        result.append((f"{title} — Desafio de Autoria", groups["authorship"]))
    if groups["sources"]:
        result.append((f"{title} — Fontes", groups["sources"]))
    return result or [(title, ["Conteúdo editorial disponível no plano."])]


def render_plan_pptx(project, plan):
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)

    cover = prs.slides.add_slide(prs.slide_layouts[6])
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = _ppt_color(BRAND.graphite)
    bar = cover.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PptInches(.16), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = _ppt_color(BRAND.primary)
    bar.line.fill.background()
    if LOGO_PATH.exists():
        cover.shapes.add_picture(str(LOGO_PATH), PptInches(10.65), PptInches(.55), width=PptInches(1.8), height=PptInches(1.8))

    _ppt_textbox(cover, BRAND.organization, .8, .65, 7.8, .4, 14, "FFFFFF", True)
    _ppt_textbox(cover, "Content Studio", .8, 1.9, 10.6, .8, 48, "FFFFFF", True)
    _ppt_textbox(cover, clean_text(project.title), .82, 2.9, 10.5, 1.05, 25, "65B5FF")
    _ppt_textbox(cover, f"{_mode_label(project)}  •  Versão {plan.version}  •  {_document_date(project)}", .82, 5.75, 10.5, .4, 15, "D0D5DD")

    number = 2
    intro = [
        f"Objetivo do projeto: {clean_text(project.objective)}",
        f"Tema: {clean_text(project.theme)}",
        f"Política de pesquisa: {_research_label(project)}",
        f"Estado editorial: {clean_text(plan.get_status_display() if hasattr(plan, 'get_status_display') else plan.status)}",
    ]
    if clean_text(project.original_intent):
        intro.append(f"Intenção original: {clean_text(project.original_intent)}")
    slide = _ppt_base(prs, "Direção editorial", number)
    _ppt_bullets(slide, intro)
    number += 1

    for title, lines in _sections(project, plan):
        is_content_block = title.startswith("Vídeo ") or " · Aula " in title
        editorial_sections = _semantic_slide_sections(title, lines) if is_content_block else [(title, lines)]
        for editorial_title, editorial_lines in editorial_sections:
            for slide_title, slide_lines in _split_for_slides(editorial_title, editorial_lines):
                slide = _ppt_base(prs, slide_title, number)
                _ppt_bullets(slide, slide_lines)
                number += 1

    sources = _source_rows(project, plan)
    source_lines = []
    for source in sources:
        text = f"[{source['kind']}] {source['title']}"
        if source["reference"]:
            text += f" — {source['reference']}"
        if source["excerpt"]:
            text += f" — {source['excerpt']}"
        source_lines.append(text)
    for slide_title, slide_lines in _split_for_slides("Fontes e proveniência", source_lines):
        slide = _ppt_base(prs, slide_title, number)
        _ppt_bullets(slide, slide_lines)
        number += 1

    props = prs.core_properties
    props.title = f"Plano editorial - {project.title}"
    props.author = BRAND.organization
    props.subject = BRAND.product
    props.keywords = "content studio, plano editorial, 3DS"

    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def render_plan_export(project, plan, export_format):
    if export_format == "pdf":
        return render_plan_pdf(project, plan)
    if export_format == "docx":
        return render_plan_docx(project, plan)
    if export_format == "pptx":
        return render_plan_pptx(project, plan)
    raise ValueError("Formato de exportação inválido.")
