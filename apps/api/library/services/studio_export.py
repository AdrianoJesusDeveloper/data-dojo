from html import escape
from io import BytesIO
import json
import re

from docx import Document
from docx.shared import Pt
from library.services.studio_scripts import spoken_text


def artifact_filename(artifact, extension):
    title = str(artifact.content.get("title") or artifact.content.get("theme") or f"artefato-{artifact.pk}")
    stem = re.sub(r"[^a-zA-Z0-9_-]+", "-", title).strip("-").lower() or "roteiro"
    return f"{stem}-{artifact.pk}.{extension}"


def render_artifact_docx(artifact):
    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading(str(artifact.content.get("title") or artifact.content.get("theme") or "Roteiro YouTube"), 0)
    document.add_paragraph(f"Status editorial: {artifact.status} · plano v{artifact.plan_version} · geração {artifact.generation}")
    for key, value in artifact.content.items():
        document.add_heading(str(key).replace("_", " ").title(), 1)
        document.add_paragraph(_text(value))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_artifact_html(artifact, teleprompter=False):
    title = str(artifact.content.get("title") or artifact.content.get("theme") or "Roteiro YouTube")
    if teleprompter:
        body = escape(spoken_text(artifact.content)).replace("\n", "<br>")
        css = "body{max-width:900px;margin:8vh auto;padding:0 5vw;background:#050505;color:#f5f5f5;font:32px/1.65 Arial}"
    else:
        body = "".join(f"<section><h2>{escape(str(key).replace('_', ' ').title())}</h2><div>{escape(_text(value)).replace(chr(10), '<br>')}</div></section>" for key, value in artifact.content.items())
        css = "body{max-width:900px;margin:40px auto;padding:0 24px;font:16px/1.6 Arial}section{margin:28px 0}"
    return f'<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>{escape(title)}</title><style>{css}</style></head><body><h1>{escape(title)}</h1><p>Status: {artifact.status}</p>{body}</body></html>'.encode("utf-8")


def _text(value):
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, indent=2)
