import io
from datetime import date

from django.utils.text import slugify
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.util import Inches as PptInches, Pt as PptPt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate, Frame, Image, KeepTogether, ListFlowable, ListItem,
    PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle,
)

from library.editorial_contracts import normalize_project_type
from professional.document_branding import BRAND as PROFESSIONAL_BRAND, LOGO_PATH


class CouncilBrand:
    organization = PROFESSIONAL_BRAND.organization
    product = "Content Studio — Conselho Editorial"
    primary = PROFESSIONAL_BRAND.primary
    graphite = PROFESSIONAL_BRAND.graphite
    accent = PROFESSIONAL_BRAND.accent
    risk = PROFESSIONAL_BRAND.risk
    muted = PROFESSIONAL_BRAND.muted
    pale_blue = PROFESSIONAL_BRAND.pale_blue
    font = PROFESSIONAL_BRAND.font


BRAND = CouncilBrand()

COUNCIL_EXPORT_MIMES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}

ROLE_LABELS = {
    "technical": "Técnico",
    "pedagogy": "Pedagogia",
    "learning_science": "Ciência da Aprendizagem",
    "technical_content": "Conteúdo Técnico",
    "youtube": "YouTube",
    "social_media": "Mídias Sociais",
    "seo": "SEO",
    "fact_checker": "Fact-checker",
}

STATUS_LABELS = {
    "queued": "Na fila",
    "running": "Em execução",
    "reviewing": "Em síntese",
    "awaiting_human_approval": "Aguardando decisão humana",
    "approved": "Aprovado",
    "revision_requested": "Revisão solicitada",
    "failed": "Falhou",
    "cancelled": "Cancelado",
}

DECISION_LABELS = {
    "approved": "APROVADO",
    "revision_requested": "REVISÃO SOLICITADA",
}


def clean_text(value):
    text = str(value or "").translate(
        str.maketrans({"\u2011": "-", "\u2010": "-", "\u00ad": "", "\u200b": ""})
    )
    return " ".join(text.split()).strip()


def _date(value):
    if value is None:
        return ""
    try:
        return value.strftime("%d/%m/%Y %H:%M")
    except AttributeError:
        return str(value)


def council_export_filename(run, extension):
    base = slugify(run.project.title) or f"projeto-{run.project_id}"
    return f"conselho-editorial-{base}-run-{run.id}.{extension}"


def _approval_snapshot(approval):
    if approval is None:
        return {
            "decision": DECISION_LABELS.get(getattr(approval, "decision", ""), "PENDENTE"),
            "notes": "",
            "decided_by": "",
            "decided_at": "",
        }
    raw_decision = clean_text(getattr(approval, "decision", ""))
    decision = "APROVADO" if raw_decision == "approved" else "REVISÃO SOLICITADA" if raw_decision == "revision" else raw_decision.upper()
    decided_by = clean_text(getattr(getattr(approval, "decided_by", None), "get_full_name", lambda: "")())
    if not decided_by:
        decided_by = clean_text(getattr(getattr(approval, "decided_by", None), "email", ""))
    return {
        "decision": decision or "PENDENTE",
        "notes": clean_text(getattr(approval, "notes", "")),
        "decided_by": decided_by,
        "decided_at": _date(getattr(approval, "created_at", None) or getattr(approval, "decided_at", None)),
    }


def build_council_snapshot(run, approval=None):
    project = run.project
    synthesis = run.final_synthesis if isinstance(run.final_synthesis, dict) else {}
    specialists = []
    for agent in run.agent_runs.all():
        output = agent.output_payload if isinstance(agent.output_payload, dict) else {}
        specialists.append({
            "role": agent.role,
            "label": ROLE_LABELS.get(agent.role, clean_text(agent.role).replace("_", " ").title()),
            "status": STATUS_LABELS.get(agent.status, clean_text(agent.status)),
            "summary": clean_text(output.get("summary")),
            "findings": [clean_text(x) for x in output.get("findings", []) if clean_text(x)],
            "recommendations": [clean_text(x) for x in output.get("recommendations", []) if clean_text(x)],
            "risks": [clean_text(x) for x in output.get("risks", []) if clean_text(x)],
        })
    approval_data = _approval_snapshot(approval)
    if run.status == "revision_requested":
        approval_data["decision"] = "REVISÃO SOLICITADA"
    elif run.status == "approved":
        approval_data["decision"] = "APROVADO"

    return {
        "title": clean_text(project.title),
        "project_type": "Conteúdo Editorial" if normalize_project_type(project.project_type) == "content" else "Formação Premium",
        "run_id": run.id,
        "plan_version": run.plan_version,
        "status": STATUS_LABELS.get(run.status, clean_text(run.status)),
        "started_at": _date(run.started_at),
        "completed_at": _date(run.completed_at),
        "summary": clean_text(synthesis.get("summary")),
        "findings": [clean_text(x) for x in synthesis.get("findings", []) if clean_text(x)],
        "recommendations": [clean_text(x) for x in synthesis.get("recommendations", []) if clean_text(x)],
        "risks": [clean_text(x) for x in synthesis.get("risks", []) if clean_text(x)],
        "specialists": specialists,
        "decision": approval_data,
    }


def _pdf_styles():
    sample = getSampleStyleSheet()
    return {
        "body": ParagraphStyle("Council Body", parent=sample["BodyText"], fontName="Helvetica", fontSize=10, leading=14, textColor=colors.HexColor("#1C1C1C"), spaceAfter=5),
        "small": ParagraphStyle("Council Small", parent=sample["BodyText"], fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor("#667085")),
        "h1": ParagraphStyle("Council H1", parent=sample["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=colors.HexColor("#0057B8"), spaceBefore=10, spaceAfter=7, keepWithNext=True),
        "h2": ParagraphStyle("Council H2", parent=sample["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=colors.HexColor("#1C1C1C"), spaceBefore=8, spaceAfter=5, keepWithNext=True),
        "coverbrand": ParagraphStyle("Council Brand", parent=sample["BodyText"], fontName="Helvetica-Bold", fontSize=11, textColor=colors.HexColor("#0057B8"), alignment=1),
        "covertitle": ParagraphStyle("Council Title", parent=sample["Title"], fontName="Helvetica-Bold", fontSize=27, leading=32, textColor=colors.HexColor("#1C1C1C"), alignment=1),
        "coverproject": ParagraphStyle("Council Project", parent=sample["Heading2"], fontName="Helvetica", fontSize=16, leading=21, textColor=colors.HexColor("#0057B8"), alignment=1),
        "decision": ParagraphStyle("Council Decision", parent=sample["Heading2"], fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=colors.HexColor("#0057B8"), alignment=1, spaceAfter=8),
    }


def _pdf_page(canvas, doc):
    if doc.page == 1:
        return
    canvas.saveState()
    width, _ = A4
    canvas.setStrokeColor(colors.HexColor("#D8E4F2"))
    canvas.line(20 * mm, 281 * mm, width - 20 * mm, 281 * mm)
    canvas.line(20 * mm, 15 * mm, width - 20 * mm, 15 * mm)
    canvas.setFont("Helvetica-Bold", 8)
    canvas.setFillColor(colors.HexColor("#0057B8"))
    canvas.drawString(20 * mm, 285 * mm, BRAND.organization)
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#667085"))
    canvas.drawString(20 * mm, 10 * mm, BRAND.product)
    canvas.drawRightString(width - 20 * mm, 10 * mm, f"Página {doc.page - 1}")
    canvas.restoreState()


def _pdf_list(story, title, items, styles):
    story.append(Paragraph(title, styles["h2"]))
    if not items:
        story.append(Paragraph("Nenhum item registrado.", styles["small"]))
        return
    story.append(ListFlowable(
        [ListItem(Paragraph(item, styles["body"]), leftIndent=5) for item in items],
        bulletType="bullet", leftIndent=16, bulletFontName="Helvetica", bulletFontSize=7,
    ))


def render_council_pdf(run, approval=None):
    s = build_council_snapshot(run, approval)
    output = io.BytesIO()
    styles = _pdf_styles()
    doc = BaseDocTemplate(
        output, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=22 * mm, bottomMargin=20 * mm,
        title=f"Conselho Editorial - {s['title']}", author=BRAND.organization,
    )
    doc.addPageTemplates(PageTemplate(
        id="Council",
        frames=(Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="body"),),
        onPage=_pdf_page,
    ))
    story = [Spacer(1, 22 * mm)]
    if LOGO_PATH.exists():
        logo = Image(str(LOGO_PATH), width=34 * mm, height=34 * mm)
        logo.hAlign = "CENTER"
        story.extend((logo, Spacer(1, 7 * mm)))
    story.extend((
        Paragraph(BRAND.organization, styles["coverbrand"]),
        Spacer(1, 8 * mm),
        Paragraph("RELATÓRIO DO CONSELHO EDITORIAL", styles["covertitle"]),
        Paragraph(s["title"], styles["coverproject"]),
        Spacer(1, 10 * mm),
    ))
    meta = [
        ("Formato", s["project_type"]),
        ("Conselho", f"#{s['run_id']}"),
        ("Plano", f"Versão {s['plan_version']}"),
        ("Estado", s["status"]),
        ("Conclusão", s["completed_at"] or "Em aberto"),
    ]
    rows = [[Paragraph(label.upper(), styles["small"]), Paragraph(value, styles["body"])] for label, value in meta]
    table = Table(rows, colWidths=(35 * mm, 90 * mm), hAlign="CENTER")
    table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 0), (-1, -1), .3, colors.HexColor("#D8E4F2")),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend((table, Spacer(1, 15 * mm), Paragraph(f"DECISÃO HUMANA: {s['decision']['decision']}", styles["decision"]), PageBreak()))

    story.append(Paragraph("Síntese executiva", styles["h1"]))
    story.append(Paragraph(s["summary"] or "Síntese não registrada.", styles["body"]))
    _pdf_list(story, "Principais constatações", s["findings"], styles)
    _pdf_list(story, "Recomendações", s["recommendations"], styles)
    _pdf_list(story, "Riscos", s["risks"], styles)

    story.append(PageBreak())
    story.append(Paragraph("Pareceres dos especialistas", styles["h1"]))
    for specialist in s["specialists"]:
        block = [Paragraph(specialist["label"], styles["h2"])]
        if specialist["summary"]:
            block.append(Paragraph(specialist["summary"], styles["body"]))
        story.append(KeepTogether(block))
        _pdf_list(story, "Constatações", specialist["findings"], styles)
        _pdf_list(story, "Recomendações", specialist["recommendations"], styles)
        _pdf_list(story, "Riscos", specialist["risks"], styles)

    story.append(PageBreak())
    story.append(Paragraph("Decisão humana e próximos passos", styles["h1"]))
    story.append(Paragraph(f"Decisão: {s['decision']['decision']}", styles["decision"]))
    if s["decision"]["notes"]:
        story.append(Paragraph(f"Observações: {s['decision']['notes']}", styles["body"]))
    if s["decision"]["decided_by"]:
        story.append(Paragraph(f"Responsável: {s['decision']['decided_by']}", styles["body"]))
    if s["decision"]["decided_at"]:
        story.append(Paragraph(f"Data da decisão: {s['decision']['decided_at']}", styles["body"]))
    next_step = (
        "Revisar o plano conforme os pareceres do Conselho antes de uma nova aprovação."
        if s["decision"]["decision"] == "REVISÃO SOLICITADA"
        else "Prosseguir somente conforme a decisão humana e o fluxo editorial vigente."
    )
    story.append(Paragraph(f"Próximo passo: {next_step}", styles["body"]))
    doc.build(story)
    return output.getvalue()


def _field_run(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _configure_docx(doc):
    colors_map = {"Title": BRAND.graphite, "Subtitle": BRAND.primary, "Heading 1": BRAND.primary, "Heading 2": BRAND.graphite, "Normal": BRAND.graphite}
    sizes = {"Title": 27, "Subtitle": 14, "Heading 1": 17, "Heading 2": 13, "Normal": 10.5}
    for name, size in sizes.items():
        style = doc.styles[name]
        style.font.name = BRAND.font
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(colors_map[name])
        style._element.rPr.rFonts.set(qn("w:ascii"), BRAND.font)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BRAND.font)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(15)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15


def _docx_header_footer(section):
    header = section.header.paragraphs[0]
    header.text = BRAND.organization
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = BRAND.font
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(BRAND.primary)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(f"{BRAND.product}  •  ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(BRAND.muted)
    _field_run(footer, "PAGE")


def _docx_bullets(doc, title, items):
    doc.add_paragraph(title, style="Heading 2")
    if not items:
        doc.add_paragraph("Nenhum item registrado.")
    else:
        for item in items:
            doc.add_paragraph(item, style="List Bullet")


def render_council_docx(run, approval=None):
    s = build_council_snapshot(run, approval)
    doc = Document()
    _configure_docx(doc)
    cover = doc.sections[0]
    cover.page_height = Inches(11.7)
    cover.page_width = Inches(8.3)
    cover.top_margin = cover.bottom_margin = Inches(.75)
    cover.left_margin = cover.right_margin = Inches(.85)
    cover.different_first_page_header_footer = True

    if LOGO_PATH.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(24)
        p.add_run().add_picture(str(LOGO_PATH), width=Inches(1.3))

    p = doc.add_paragraph(BRAND.organization)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.color.rgb = RGBColor.from_string(BRAND.primary)

    p = doc.add_paragraph("RELATÓRIO DO CONSELHO EDITORIAL", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(32)
    p = doc.add_paragraph(s["title"], style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = [
        ("Formato", s["project_type"]), ("Conselho", f"#{s['run_id']}"),
        ("Plano", f"Versão {s['plan_version']}"), ("Estado", s["status"]),
        ("Conclusão", s["completed_at"] or "Em aberto"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row, (label, value) in zip(table.rows, meta):
        row.cells[0].text = label.upper()
        row.cells[1].text = value
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = doc.add_paragraph(f"DECISÃO HUMANA: {s['decision']['decision']}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(15)
    p.runs[0].font.color.rgb = RGBColor.from_string(BRAND.primary)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_height = Inches(11.7)
    section.page_width = Inches(8.3)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    _docx_header_footer(section)

    doc.add_paragraph("Síntese executiva", style="Heading 1")
    doc.add_paragraph(s["summary"] or "Síntese não registrada.")
    _docx_bullets(doc, "Principais constatações", s["findings"])
    _docx_bullets(doc, "Recomendações", s["recommendations"])
    _docx_bullets(doc, "Riscos", s["risks"])

    doc.add_page_break()
    doc.add_paragraph("Pareceres dos especialistas", style="Heading 1")
    for specialist in s["specialists"]:
        doc.add_paragraph(specialist["label"], style="Heading 2")
        if specialist["summary"]:
            doc.add_paragraph(specialist["summary"])
        _docx_bullets(doc, "Constatações", specialist["findings"])
        _docx_bullets(doc, "Recomendações", specialist["recommendations"])
        _docx_bullets(doc, "Riscos", specialist["risks"])

    doc.add_page_break()
    doc.add_paragraph("Decisão humana e próximos passos", style="Heading 1")
    doc.add_paragraph(f"Decisão: {s['decision']['decision']}")
    if s["decision"]["notes"]:
        doc.add_paragraph(f"Observações: {s['decision']['notes']}")
    if s["decision"]["decided_by"]:
        doc.add_paragraph(f"Responsável: {s['decision']['decided_by']}")
    if s["decision"]["decided_at"]:
        doc.add_paragraph(f"Data da decisão: {s['decision']['decided_at']}")
    next_step = "Revisar o plano conforme os pareceres do Conselho antes de uma nova aprovação." if s["decision"]["decision"] == "REVISÃO SOLICITADA" else "Prosseguir somente conforme a decisão humana e o fluxo editorial vigente."
    doc.add_paragraph(f"Próximo passo: {next_step}")

    props = doc.core_properties
    props.title = f"Conselho Editorial - {s['title']}"
    props.author = BRAND.organization
    props.subject = BRAND.product
    props.keywords = "content studio, conselho editorial, decisão humana, 3DS"
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _ppt_color(value):
    return PptRGBColor.from_string(value)


def _ppt_textbox(slide, text, x, y, w, h, size=20, color=None, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    p = frame.paragraphs[0]
    p.text = clean_text(text)
    p.alignment = align
    p.font.name = BRAND.font
    p.font.size = PptPt(size)
    p.font.bold = bold
    p.font.color.rgb = _ppt_color(color or BRAND.graphite)
    return shape


def _ppt_footer(slide, number):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(.55), PptInches(7.12), PptInches(12.23), PptInches(.02))
    line.fill.solid()
    line.fill.fore_color.rgb = _ppt_color("D8E4F2")
    line.line.fill.background()
    _ppt_textbox(slide, BRAND.organization, .58, 7.18, 5.8, .22, 8, BRAND.muted)
    _ppt_textbox(slide, str(number), 12.1, 7.18, .65, .22, 8, BRAND.muted, align=PP_ALIGN.RIGHT)


def _ppt_base(prs, title, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _ppt_color("FFFFFF")
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(.55), PptInches(.55), PptInches(.1), PptInches(.64))
    accent.fill.solid()
    accent.fill.fore_color.rgb = _ppt_color(BRAND.primary)
    accent.line.fill.background()
    _ppt_textbox(slide, title, .82, .5, 11.3, .75, 30, BRAND.graphite, True)
    _ppt_footer(slide, number)
    return slide


def _ppt_bullets(slide, items, max_items=6):
    items = [clean_text(item) for item in items if clean_text(item)][:max_items]
    if not items:
        items = ["Nenhum item registrado."]
    shape = slide.shapes.add_textbox(PptInches(.85), PptInches(1.5), PptInches(11.6), PptInches(5.3))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, item in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = f"• {item}"
        p.font.name = BRAND.font
        p.font.size = PptPt(18 if len(items) > 4 else 20)
        p.font.color.rgb = _ppt_color(BRAND.graphite)
        p.space_after = PptPt(10)


def _slide_chunks(title, items, size=6):
    items = list(items or [])
    if not items:
        return [(title, [])]
    return [
        (title if i == 0 else f"{title} — continuação", items[i:i + size])
        for i in range(0, len(items), size)
    ]


def render_council_pptx(run, approval=None):
    s = build_council_snapshot(run, approval)
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)

    cover = prs.slides.add_slide(prs.slide_layouts[6])
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = _ppt_color(BRAND.graphite)
    bar = cover.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PptInches(.16), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = _ppt_color(BRAND.primary)
    bar.line.fill.background()
    if LOGO_PATH.exists():
        cover.shapes.add_picture(str(LOGO_PATH), PptInches(10.65), PptInches(.55), width=PptInches(1.8), height=PptInches(1.8))
    _ppt_textbox(cover, BRAND.organization, .8, .65, 7.8, .4, 14, "FFFFFF", True)
    _ppt_textbox(cover, "Conselho Editorial", .8, 1.9, 10.6, .8, 46, "FFFFFF", True)
    _ppt_textbox(cover, s["title"], .82, 2.9, 10.5, 1.05, 25, "65B5FF")
    _ppt_textbox(cover, f"Run #{s['run_id']}  •  Plano v{s['plan_version']}  •  {s['decision']['decision']}", .82, 5.75, 10.8, .4, 15, "D0D5DD")

    number = 2
    slide = _ppt_base(prs, "Visão geral", number)
    _ppt_bullets(slide, [
        f"Projeto: {s['title']}",
        f"Formato: {s['project_type']}",
        f"Plano analisado: versão {s['plan_version']}",
        f"Estado do Conselho: {s['status']}",
        f"Decisão humana: {s['decision']['decision']}",
        f"Conclusão: {s['completed_at'] or 'Em aberto'}",
    ])
    number += 1

    slide = _ppt_base(prs, "Síntese executiva", number)
    _ppt_bullets(slide, [s["summary"]])
    number += 1

    for title, items in (
        ("Principais constatações", s["findings"]),
        ("Recomendações prioritárias", s["recommendations"]),
        ("Riscos editoriais", s["risks"]),
    ):
        for slide_title, chunk in _slide_chunks(title, items):
            slide = _ppt_base(prs, slide_title, number)
            _ppt_bullets(slide, chunk)
            number += 1

    for specialist in s["specialists"]:
        items = []
        if specialist["summary"]:
            items.append(f"Síntese: {specialist['summary']}")
        items.extend(f"Constatação: {x}" for x in specialist["findings"][:2])
        items.extend(f"Recomendação: {x}" for x in specialist["recommendations"][:2])
        items.extend(f"Risco: {x}" for x in specialist["risks"][:2])
        for slide_title, chunk in _slide_chunks(f"Parecer — {specialist['label']}", items):
            slide = _ppt_base(prs, slide_title, number)
            _ppt_bullets(slide, chunk)
            number += 1

    slide = _ppt_base(prs, "Decisão humana", number)
    decision_lines = [f"Decisão: {s['decision']['decision']}"]
    if s["decision"]["notes"]:
        decision_lines.append(f"Observações: {s['decision']['notes']}")
    if s["decision"]["decided_by"]:
        decision_lines.append(f"Responsável: {s['decision']['decided_by']}")
    if s["decision"]["decided_at"]:
        decision_lines.append(f"Data: {s['decision']['decided_at']}")
    _ppt_bullets(slide, decision_lines)
    number += 1

    slide = _ppt_base(prs, "Próximos passos", number)
    _ppt_bullets(slide, [
        "Revisar o plano conforme os pareceres do Conselho." if s["decision"]["decision"] == "REVISÃO SOLICITADA" else "Prosseguir conforme a decisão humana registrada.",
        "Preservar a rastreabilidade entre versão do plano, pareceres e decisão.",
        "Submeter mudanças relevantes a nova revisão editorial antes da publicação.",
    ])

    props = prs.core_properties
    props.title = f"Conselho Editorial - {s['title']}"
    props.author = BRAND.organization
    props.subject = BRAND.product
    props.keywords = "content studio, conselho editorial, decisão humana, 3DS"
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def render_council_export(run, approval, export_format):
    if export_format == "pdf":
        return render_council_pdf(run, approval)
    if export_format == "docx":
        return render_council_docx(run, approval)
    if export_format == "pptx":
        return render_council_pptx(run, approval)
    raise ValueError("Formato de exportação inválido.")
