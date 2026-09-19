import io
import re
from html import escape
from xml.sax.saxutils import escape as xml_escape

from django.utils.text import slugify
from docx import Document
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.util import Inches, Pt as PptPt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer

from professional.document_branding import BRAND as PROFESSIONAL_BRAND


SECTION_EXPORT_MIMES = {
    "html": "text/html; charset=utf-8",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}

SECTION_LABELS = {
    "research": "Pesquisa Fundamentada",
    "dossier": "Dossiê",
    "editorial_view": "Visão Editorial",
    "modules": "Módulos",
    "exercises": "Exercícios",
    "council": "Conselho Editorial",
    "editorial_content": "Conteúdo Editorial",
    "artifacts": "Artefatos Editoriais",
    "generations": "Histórico de Gerações",
}


def _clean(value):
    return " ".join(str(value or "").replace("\u00ad", "").replace("\u200b", "").split()).strip()


def _label(value):
    return _clean(value).replace("_", " ").strip().capitalize()


def _serialize_evidence(item):
    return {
        "tipo": getattr(item, "source_kind", ""),
        "título": getattr(item, "title", ""),
        "domínio": getattr(item, "domain", ""),
        "url": getattr(item, "url", ""),
        "trecho": getattr(item, "excerpt", ""),
        "consultado_em": getattr(item, "retrieved_at", None),
    }


def _plan_architecture(project):
    plan = getattr(project, "modernization_plan", None)
    return (getattr(plan, "proposed_architecture", None) or {}) if plan else {}


def build_section_snapshot(project, section):
    if section not in SECTION_LABELS:
        raise ValueError("Seção de exportação não suportada.")

    payload = None

    if section == "research":
        try:
            context = project.research_context
        except Exception:
            context = None
        if context is None:
            payload = {"status": "Pesquisa ainda não executada."}
        else:
            payload = {
                "status": context.status,
                "política": context.policy,
                "dossiê_da_pesquisa": context.dossier,
                "evidências": [_serialize_evidence(item) for item in context.evidence.all()],
            }

    elif section == "dossier":
        versions = list(project.dossier_versions.all().order_by("-version"))
        payload = {
            "versão_atual": (
                {
                    "versão": versions[0].version,
                    "status": versions[0].status,
                    "origem": versions[0].origin,
                    "política": versions[0].research_policy,
                    "conteúdo": versions[0].content,
                    "referências": versions[0].references_snapshot,
                    "criado_em": versions[0].created_at,
                }
                if versions
                else {"status": "Nenhuma versão de Dossiê criada."}
            ),
            "histórico": [
                {
                    "versão": item.version,
                    "status": item.status,
                    "origem": item.origin,
                    "criado_em": item.created_at,
                }
                for item in versions
            ],
        }

    elif section == "editorial_view":
        payload = _plan_architecture(project) or {"status": "Plano editorial ainda não gerado."}

    elif section == "modules":
        architecture = _plan_architecture(project)
        payload = {
            "módulos": architecture.get("modules", []),
        }

    elif section == "exercises":
        architecture = _plan_architecture(project)
        groups = []
        for module_index, module in enumerate(architecture.get("modules", []) or [], start=1):
            if not isinstance(module, dict):
                continue
            lessons = []
            for lesson_index, lesson in enumerate(module.get("lessons", []) or [], start=1):
                if not isinstance(lesson, dict):
                    continue
                lessons.append({
                    "aula": lesson.get("title") or f"Aula {lesson_index}",
                    "exercício": lesson.get("exercise"),
                    "desafio_sem_ia": lesson.get("without_ai_challenge"),
                    "desafio_de_autoria": lesson.get("authorship_challenge"),
                })
            groups.append({
                "módulo": module.get("title") or f"Módulo {module_index}",
                "exercícios": module.get("exercises"),
                "kata": module.get("kata"),
                "projeto_prático": module.get("practical_project"),
                "avaliação": module.get("assessment"),
                "aulas": lessons,
            })
        payload = {"exercícios": groups}

    elif section == "council":
        run = project.council_runs.prefetch_related("agent_runs").order_by("-id").first()
        if run is None:
            payload = {"status": "Nenhuma execução do Conselho Editorial."}
        else:
            from library.services.council_export import build_council_snapshot
            payload = build_council_snapshot(run)

    elif section == "editorial_content":
        try:
            package = project.content_package
        except Exception:
            package = None
        payload = {
            "itens_gerados": (package.generated_items if package else []),
            "status_de_publicação": getattr(package, "publication_status", "") if package else "",
        }

    elif section == "artifacts":
        payload = {
            "artefatos": [
                {
                    "id": item.id,
                    "tipo": item.artifact_type,
                    "alvo": item.target_type,
                    "status": item.status,
                    "versão_do_plano": item.plan_version,
                    "geração": item.generation,
                    "conteúdo": item.content,
                    "criado_em": item.created_at,
                    "atualizado_em": item.updated_at,
                }
                for item in project.artifacts.all().order_by("-id")
            ]
        }

    elif section == "generations":
        try:
            package = project.content_package
        except Exception:
            package = None
        items = (package.generated_items if package else []) or []
        payload = {"gerações": items}

    return {
        "title": SECTION_LABELS[section],
        "project": project.title,
        "objective": project.objective,
        "section": section,
        "content": payload,
    }


def section_export_filename(project, section, extension):
    project_slug = slugify(project.title) or f"projeto-{project.pk}"
    section_slug = slugify(SECTION_LABELS.get(section, section)) or section
    return f"content-studio-{project_slug}-{section_slug}.{extension}"


def _html_value(value):
    if value is None or value == "":
        return '<p class="muted">Não informado.</p>'
    if isinstance(value, dict):
        blocks = []
        for key, item in value.items():
            blocks.append(
                f'<section class="block"><h3>{escape(_label(key))}</h3>{_html_value(item)}</section>'
            )
        return "".join(blocks)
    if isinstance(value, (list, tuple)):
        if not value:
            return '<p class="muted">Nenhum item registrado.</p>'
        return "<div class=\"list\">" + "".join(
            f'<article class="item">{_html_value(item)}</article>' for item in value
        ) + "</div>"
    return f"<p>{escape(_clean(value))}</p>"


def render_section_html(snapshot):
    body = _html_value(snapshot["content"])
    return f"""<!doctype html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<title>{escape(snapshot["title"])} — {escape(snapshot["project"])}</title>
<style>
body{{font-family:Arial,sans-serif;color:#1c1c1c;max-width:980px;margin:40px auto;padding:0 28px;line-height:1.55}}
header{{border-bottom:2px solid #0057b8;padding-bottom:18px;margin-bottom:24px}}
h1{{margin:0;color:#0057b8}} h2{{margin:8px 0 0}} h3{{font-size:14px;text-transform:uppercase;letter-spacing:.04em;margin:0 0 8px}}
.block{{border:1px solid #d8e4f2;border-radius:10px;padding:14px;margin:12px 0}}
.item{{border-left:3px solid #0057b8;padding:10px 14px;margin:10px 0;background:#f7f9fc}}
.muted{{color:#667085}}
@media print{{body{{margin:0;max-width:none}}}}
</style>
</head>
<body>
<header><h1>{escape(snapshot["title"])}</h1><h2>{escape(snapshot["project"])}</h2><p>{escape(snapshot["objective"])}</p></header>
{body}
</body>
</html>""".encode("utf-8")


def _walk(value, level=0, heading=None):
    if heading:
        yield ("heading", level, _label(heading))
    if value is None or value == "":
        yield ("text", level, "Não informado.")
    elif isinstance(value, dict):
        if not value:
            yield ("text", level, "Nenhum item registrado.")
        for key, item in value.items():
            yield from _walk(item, level + 1, key)
    elif isinstance(value, (list, tuple)):
        if not value:
            yield ("text", level, "Nenhum item registrado.")
        for index, item in enumerate(value, start=1):
            if isinstance(item, (dict, list, tuple)):
                yield ("heading", level + 1, f"Item {index}")
                yield from _walk(item, level + 2)
            else:
                yield ("bullet", level + 1, _clean(item))
    else:
        yield ("text", level, _clean(value))


def render_section_pdf(snapshot):
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "SectionTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=22, leading=27, textColor=colors.HexColor("#0057B8"), spaceAfter=8,
    )
    subtitle_style = ParagraphStyle(
        "SectionSubtitle", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=13, leading=17, textColor=colors.HexColor("#1C1C1C"), spaceAfter=10,
    )
    body_style = ParagraphStyle(
        "SectionBody", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=9.5, leading=13.5, textColor=colors.HexColor("#1C1C1C"), spaceAfter=5,
    )
    small_style = ParagraphStyle(
        "SectionSmall", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=8.5, leading=11, textColor=colors.HexColor("#667085"), spaceAfter=5,
    )
    heading_styles = {
        1: ParagraphStyle("H1", parent=styles["Heading1"], fontSize=14, leading=18, textColor=colors.HexColor("#0057B8"), spaceBefore=8, spaceAfter=5),
        2: ParagraphStyle("H2", parent=styles["Heading2"], fontSize=12, leading=16, textColor=colors.HexColor("#1C1C1C"), spaceBefore=6, spaceAfter=4),
        3: ParagraphStyle("H3", parent=styles["Heading3"], fontSize=10.5, leading=14, textColor=colors.HexColor("#344054"), spaceBefore=5, spaceAfter=3),
    }

    doc = BaseDocTemplate(
        output, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
        title=f'{snapshot["title"]} - {snapshot["project"]}',
        author=PROFESSIONAL_BRAND.organization,
    )
    doc.addPageTemplates(PageTemplate(id="content", frames=(Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="body"),)))
    story = [
        Paragraph(xml_escape(snapshot["title"]), title_style),
        Paragraph(xml_escape(snapshot["project"]), subtitle_style),
        Paragraph(xml_escape(snapshot["objective"]), small_style),
        Spacer(1, 4 * mm),
    ]

    for kind, level, value in _walk(snapshot["content"]):
        safe = xml_escape(value)
        if kind == "heading":
            story.append(Paragraph(safe, heading_styles.get(min(max(level, 1), 3), heading_styles[3])))
        elif kind == "bullet":
            story.append(Paragraph(f"• {safe}", body_style))
        else:
            story.append(Paragraph(safe, body_style))

    doc.build(story)
    return output.getvalue()


def render_section_docx(snapshot):
    doc = Document()
    doc.core_properties.title = f'{snapshot["title"]} - {snapshot["project"]}'
    doc.core_properties.author = PROFESSIONAL_BRAND.organization

    title = doc.add_heading(snapshot["title"], 0)
    title.runs[0].font.color.rgb = RGBColor(0x00, 0x57, 0xB8)
    doc.add_heading(snapshot["project"], level=1)
    if snapshot["objective"]:
        p = doc.add_paragraph(snapshot["objective"])
        p.style = doc.styles["Subtitle"]

    for kind, level, value in _walk(snapshot["content"]):
        if kind == "heading":
            doc.add_heading(value, level=min(max(level, 1), 3))
        elif kind == "bullet":
            doc.add_paragraph(value, style="List Bullet")
        else:
            doc.add_paragraph(value)

    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = PROFESSIONAL_BRAND.font
        if style_name == "Normal":
            style.font.size = Pt(10.5)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _ppt_split_text(value, max_chars=150):
    """Split long editorial text into readable presentation blocks."""
    value = _clean(value)
    if not value:
        return []

    sentence_parts = [
        item.strip()
        for item in re.split(r"(?<=[.!?;:])\s+|(?:\s+[•·]\s+)", value)
        if item.strip()
    ]

    result = []
    for part in sentence_parts or [value]:
        if len(part) <= max_chars:
            result.append(part)
            continue

        words = part.split()
        current = []
        current_len = 0
        for word in words:
            projected = current_len + len(word) + (1 if current else 0)
            if current and projected > max_chars:
                result.append(" ".join(current))
                current = [word]
                current_len = len(word)
            else:
                current.append(word)
                current_len = projected
        if current:
            result.append(" ".join(current))

    return result


def _ppt_add_chunk(chunks, title, lines):
    """Split content into compact slides to avoid overflowing placeholders."""
    clean_lines = []
    for line in lines:
        if not line:
            continue
        line = str(line)
        is_bullet = line.startswith("• ")
        raw = line[2:] if is_bullet else line
        for part in _ppt_split_text(raw):
            clean_lines.append(("• " if is_bullet else "") + part)

    if not clean_lines:
        chunks.append((title, ["Nenhum item registrado."]))
        return

    page = []
    page_chars = 0
    continuation = 1

    for line in clean_lines:
        # Conservative limits keep the text readable at presentation size.
        if page and (len(page) >= 5 or page_chars + len(line) > 620):
            slide_title = title if continuation == 1 else f"{title} — continuação {continuation}"
            chunks.append((slide_title, page))
            continuation += 1
            page = []
            page_chars = 0

        page.append(line)
        page_chars += len(line)

    if page:
        slide_title = title if continuation == 1 else f"{title} — continuação {continuation}"
        chunks.append((slide_title, page))


def _ppt_chunks(snapshot):
    chunks = []

    # Long project objectives no longer go on the cover.
    if snapshot.get("objective"):
        _ppt_add_chunk(chunks, "Objetivo do projeto", [snapshot["objective"]])

    current_title = snapshot["title"]
    lines = []

    for kind, level, value in _walk(snapshot["content"]):
        if kind == "heading" and level <= 2:
            if lines:
                _ppt_add_chunk(chunks, current_title, lines)
                lines = []
            current_title = value
            continue

        prefix = "• " if kind == "bullet" else ""
        lines.append(prefix + value)

    if lines:
        _ppt_add_chunk(chunks, current_title, lines)

    if not chunks:
        chunks.append((snapshot["title"], ["Nenhum item registrado."]))

    return chunks



PPT_FIELD_LABELS = {
    "theme": "Tema",
    "tema": "Tema",
    "thesis": "Tese",
    "tese": "Tese",
    "problem": "Problema",
    "problema": "Problema",
    "audience": "Público",
    "publico": "Público",
    "promise": "Promessa",
    "promessa": "Promessa",
    "objectives": "Objetivos",
    "learning_objectives": "Objetivos de aprendizagem",
    "objetivos_de_aprendizagem": "Objetivos de aprendizagem",
    "competencies": "Competências",
    "competencias": "Competências",
    "prerequisites": "Pré-requisitos",
    "pre_requisites": "Pré-requisitos",
    "pre_requisitos": "Pré-requisitos",
    "practices": "Práticas",
    "practice": "Práticas",
    "exercises": "Exercícios",
    "expected_evidence": "Evidências esperadas",
    "evidencias_esperadas": "Evidências esperadas",
    "criteria": "Critérios",
    "assessment_criteria": "Critérios de avaliação",
    "editorial_insights": "Insights editoriais",
    "insights": "Insights",
    "limitations": "Limitações",
    "limitacoes": "Limitações",
    "counterpoints": "Contrapontos",
    "risks": "Riscos",
    "sources": "Fontes",
    "recommended_sources": "Fontes recomendadas",
    "references": "Referências",
    "referencias": "Referências",
    "coverage": "Cobertura",
    "source_coverage": "Cobertura de fontes",
    "notes": "Anotações",
}


def _ppt_field_label(value):
    key = _clean(value).lower()
    return PPT_FIELD_LABELS.get(key, _label(value))


def _ppt_trim(value, max_chars):
    value = _clean(value)
    if len(value) <= max_chars:
        return value
    cut = value[: max_chars - 1].rsplit(" ", 1)[0].strip()
    return f"{cut}…"


def _ppt_compact_lines(value, *, max_lines=6, max_total_chars=920, prefix=""):
    """Create a concise, faithful presentation view without replacing the source content."""
    lines = []
    used_chars = 0
    truncated = False

    def add_line(line):
        nonlocal used_chars, truncated
        line = _clean(line)
        if not line:
            return
        if len(lines) >= max_lines:
            truncated = True
            return

        remaining = max_total_chars - used_chars
        if remaining <= 0:
            truncated = True
            return

        if len(line) > remaining:
            line = _ppt_trim(line, max(80, remaining))
            truncated = True

        lines.append(line)
        used_chars += len(line)

    def visit(item, label_prefix=""):
        nonlocal truncated
        if len(lines) >= max_lines or used_chars >= max_total_chars:
            truncated = True
            return

        if item is None or item == "":
            return

        if isinstance(item, dict):
            for key, child in item.items():
                child_label = _ppt_field_label(key)
                next_prefix = f"{label_prefix} · {child_label}" if label_prefix else child_label

                if isinstance(child, (dict, list, tuple)):
                    visit(child, next_prefix)
                else:
                    add_line(f"{next_prefix}: {_clean(child)}")
            return

        if isinstance(item, (list, tuple)):
            for child in item:
                if isinstance(child, (dict, list, tuple)):
                    visit(child, label_prefix)
                else:
                    text = _clean(child)
                    add_line(f"{label_prefix}: {text}" if label_prefix else text)
            return

        add_line(f"{label_prefix}: {_clean(item)}" if label_prefix else _clean(item))

    visit(value, prefix)

    if not lines:
        lines = ["Nenhum item registrado."]

    if truncated:
        note = "Conteúdo completo disponível nas exportações HTML, PDF e DOCX."
        if len(lines) < max_lines and used_chars + len(note) <= max_total_chars:
            lines.append(note)
        else:
            lines[-1] = _ppt_trim(lines[-1], max(90, len(lines[-1]) - len(note) // 2))

    return lines


def _ppt_research_chunks(snapshot):
    """Executive presentation for Pesquisa Fundamentada.

    The detailed research remains intact in HTML/PDF/DOCX; PPTX is intentionally
    an executive presentation rather than a verbatim document dump.
    """
    content = snapshot.get("content") or {}
    chunks = []

    if snapshot.get("objective"):
        chunks.append(
            (
                "Objetivo do projeto",
                _ppt_compact_lines(snapshot["objective"], max_lines=6, max_total_chars=1000),
            )
        )

    overview_lines = []
    if content.get("status"):
        overview_lines.append(f"Status: {_clean(content.get('status'))}")
    if content.get("política"):
        overview_lines.append(f"Política de pesquisa: {_clean(content.get('política'))}")

    evidences = content.get("evidências") or []
    if isinstance(evidences, (list, tuple)):
        overview_lines.append(f"Evidências registradas: {len(evidences)}")

        acervo = sum(
            1
            for item in evidences
            if isinstance(item, dict)
            and _clean(item.get("tipo")).lower() not in {"web", "website", "url"}
        )
        web = sum(
            1
            for item in evidences
            if isinstance(item, dict)
            and _clean(item.get("tipo")).lower() in {"web", "website", "url"}
        )
        if evidences:
            overview_lines.append(f"Cobertura: {acervo} do acervo · {web} da web")

    if overview_lines:
        chunks.append(("Status e cobertura", overview_lines))

    dossier = content.get("dossiê_da_pesquisa")
    if isinstance(dossier, dict) and dossier:
        preferred_order = [
            "theme",
            "tema",
            "problem",
            "problema",
            "thesis",
            "tese",
            "audience",
            "publico",
            "promise",
            "promessa",
            "objectives",
            "learning_objectives",
            "objetivos_de_aprendizagem",
            "competencies",
            "competencias",
            "prerequisites",
            "pre_requisitos",
            "practices",
            "practice",
            "exercises",
            "expected_evidence",
            "evidencias_esperadas",
            "criteria",
            "assessment_criteria",
            "editorial_insights",
            "insights",
            "limitations",
            "limitacoes",
            "counterpoints",
            "risks",
            "recommended_sources",
            "sources",
            "references",
            "referencias",
        ]

        selected_keys = []
        for key in preferred_order:
            if key in dossier and key not in selected_keys:
                selected_keys.append(key)

        for key in dossier.keys():
            if key not in selected_keys:
                selected_keys.append(key)

        max_dossier_slides = 12
        shown_keys = selected_keys[:max_dossier_slides]

        for key in shown_keys:
            chunks.append(
                (
                    _ppt_field_label(key),
                    _ppt_compact_lines(
                        dossier.get(key),
                        max_lines=7,
                        max_total_chars=1000,
                    ),
                )
            )

        remaining_keys = selected_keys[max_dossier_slides:]
        if remaining_keys:
            chunks.append(
                (
                    "Outros pontos do dossiê",
                    [
                        "Campos adicionais preservados no relatório completo:",
                        *[_ppt_field_label(key) for key in remaining_keys[:8]],
                        "Consulte HTML, PDF ou DOCX para o conteúdo integral.",
                    ],
                )
            )
    elif dossier:
        chunks.append(
            (
                "Dossiê da pesquisa",
                _ppt_compact_lines(dossier, max_lines=7, max_total_chars=1000),
            )
        )

    if isinstance(evidences, (list, tuple)) and evidences:
        chunks.append(
            (
                "Evidências e fontes",
                [
                    f"Total de evidências: {len(evidences)}",
                    "A apresentação mostra uma seleção das evidências para manter legibilidade.",
                    "O conjunto completo permanece disponível nas exportações HTML, PDF e DOCX.",
                ],
            )
        )

        max_evidence_slides = 5
        for index, evidence in enumerate(evidences[:max_evidence_slides], start=1):
            if not isinstance(evidence, dict):
                lines = _ppt_compact_lines(evidence, max_lines=5, max_total_chars=850)
                title = f"Evidência {index}"
            else:
                title_text = (
                    evidence.get("título")
                    or evidence.get("title")
                    or evidence.get("domínio")
                    or evidence.get("domain")
                    or f"Evidência {index}"
                )
                title = f"Evidência {index} — {_ppt_trim(title_text, 72)}"
                reduced = {
                    "tipo": evidence.get("tipo"),
                    "domínio": evidence.get("domínio"),
                    "url": evidence.get("url"),
                    "trecho": evidence.get("trecho"),
                    "consultado_em": evidence.get("consultado_em"),
                }
                lines = _ppt_compact_lines(reduced, max_lines=6, max_total_chars=900)

            chunks.append((title, lines))

        if len(evidences) > max_evidence_slides:
            chunks.append(
                (
                    "Demais evidências",
                    [
                        f"{len(evidences) - max_evidence_slides} evidências adicionais não foram reproduzidas em slides.",
                        "Todas permanecem disponíveis no HTML, PDF e DOCX.",
                    ],
                )
            )

    if not chunks:
        chunks.append((snapshot["title"], ["Nenhum item registrado."]))

    return chunks


# --- Identidade visual Data Driven Dojô para PPTX ----------------------------

PPT_BG = PptRGBColor(0x0B, 0x0D, 0x10)
PPT_PANEL = PptRGBColor(0x15, 0x18, 0x1D)
PPT_PANEL_2 = PptRGBColor(0x1C, 0x20, 0x27)
PPT_TEXT = PptRGBColor(0xF5, 0xF5, 0xF2)
PPT_MUTED = PptRGBColor(0xA4, 0xAA, 0xB4)
PPT_RED = PptRGBColor(0xD7, 0x19, 0x20)
PPT_RED_SOFT = PptRGBColor(0x7D, 0x12, 0x16)
PPT_BLUE = PptRGBColor(0x4A, 0xA3, 0xFF)
PPT_AMBER = PptRGBColor(0xF5, 0xB9, 0x42)
PPT_WHITE = PptRGBColor(0xFF, 0xFF, 0xFF)


def _ppt_set_background(slide, color=PPT_BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _ppt_add_brand_bar(slide, *, accent=PPT_RED):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0,
        0,
        Inches(0.16),
        Inches(7.5),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    top_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.16),
        0,
        Inches(13.173),
        Inches(0.05),
    )
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = accent
    top_line.line.fill.background()


def _ppt_add_footer(slide, section_title):
    footer = slide.shapes.add_textbox(
        Inches(0.55),
        Inches(7.06),
        Inches(12.1),
        Inches(0.22),
    )
    frame = footer.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = f"DATA DRIVEN DOJÔ  •  CONTENT STUDIO  •  {section_title}"
    p.font.name = PROFESSIONAL_BRAND.font
    p.font.size = PptPt(8)
    p.font.color.rgb = PPT_MUTED
    p.alignment = PP_ALIGN.LEFT


def _ppt_add_watermark(slide):
    mark = slide.shapes.add_textbox(
        Inches(10.55),
        Inches(0.45),
        Inches(2.1),
        Inches(1.15),
    )
    frame = mark.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = "3DS"
    p.font.name = PROFESSIONAL_BRAND.font
    p.font.size = PptPt(44)
    p.font.bold = True
    p.font.color.rgb = PptRGBColor(0x22, 0x25, 0x2B)
    p.alignment = PP_ALIGN.RIGHT


def _ppt_style_title(shape, *, cover=False):
    frame = shape.text_frame
    frame.word_wrap = True
    for paragraph in frame.paragraphs:
        paragraph.font.name = PROFESSIONAL_BRAND.font
        paragraph.font.size = PptPt(30 if not cover else 34)
        paragraph.font.bold = True
        paragraph.font.color.rgb = PPT_WHITE


def _ppt_style_body(frame):
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.05)
    frame.margin_bottom = Inches(0.05)

    for paragraph in frame.paragraphs:
        paragraph.font.name = PROFESSIONAL_BRAND.font
        paragraph.font.size = PptPt(16)
        paragraph.font.color.rgb = PPT_TEXT
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = PptPt(7)


def _ppt_add_content_panel(slide):
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.62),
        Inches(1.55),
        Inches(12.05),
        Inches(5.1),
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = PPT_PANEL
    panel.line.color.rgb = PptRGBColor(0x2C, 0x32, 0x3B)
    panel.line.width = PptPt(0.8)
    return panel


def _ppt_move_title(shape):
    shape.left = Inches(0.72)
    shape.top = Inches(0.48)
    shape.width = Inches(10.5)
    shape.height = Inches(0.82)


def _ppt_move_body(shape):
    shape.left = Inches(0.88)
    shape.top = Inches(1.85)
    shape.width = Inches(11.35)
    shape.height = Inches(4.55)


def _ppt_cover(slide, snapshot):
    _ppt_set_background(slide, PPT_BG)
    _ppt_add_brand_bar(slide, accent=PPT_RED)
    _ppt_add_watermark(slide)

    # Decorative red block
    block = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.85),
        Inches(1.35),
        Inches(0.24),
        Inches(3.55),
    )
    block.fill.solid()
    block.fill.fore_color.rgb = PPT_RED
    block.line.fill.background()

    title_shape = slide.shapes.title
    title_shape.left = Inches(1.35)
    title_shape.top = Inches(1.55)
    title_shape.width = Inches(9.8)
    title_shape.height = Inches(1.5)
    title_shape.text = snapshot["title"]
    _ppt_style_title(title_shape, cover=True)

    subtitle = slide.placeholders[1]
    subtitle.left = Inches(1.38)
    subtitle.top = Inches(3.25)
    subtitle.width = Inches(9.7)
    subtitle.height = Inches(1.4)
    subtitle.text = snapshot["project"]
    subtitle.text_frame.word_wrap = True
    for paragraph in subtitle.text_frame.paragraphs:
        paragraph.font.name = PROFESSIONAL_BRAND.font
        paragraph.font.size = PptPt(22)
        paragraph.font.color.rgb = PPT_MUTED
        paragraph.alignment = PP_ALIGN.LEFT

    tag = slide.shapes.add_textbox(
        Inches(1.38),
        Inches(5.35),
        Inches(5.5),
        Inches(0.5),
    )
    p = tag.text_frame.paragraphs[0]
    p.text = "DETERMINAÇÃO  •  DISCIPLINA  •  DIREÇÃO"
    p.font.name = PROFESSIONAL_BRAND.font
    p.font.size = PptPt(10)
    p.font.bold = True
    p.font.color.rgb = PPT_RED

    _ppt_add_footer(slide, snapshot["title"])


def _ppt_apply_content_theme(slide, section_title):
    _ppt_set_background(slide, PPT_BG)
    _ppt_add_brand_bar(slide, accent=PPT_RED)
    _ppt_add_watermark(slide)
    _ppt_add_footer(slide, section_title)

    title_shape = slide.shapes.title
    _ppt_move_title(title_shape)
    _ppt_style_title(title_shape)

    body_shape = slide.placeholders[1]
    _ppt_move_body(body_shape)

    panel = _ppt_add_content_panel(slide)
    # ensure text stays above the decorative panel
    slide.shapes._spTree.remove(body_shape._element)
    slide.shapes._spTree.append(body_shape._element)

    return body_shape.text_frame


def render_section_pptx(snapshot):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Capa personalizada Data Driven Dojô
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    _ppt_cover(slide, snapshot)

    chunks = (
        _ppt_research_chunks(snapshot)
        if snapshot.get("section") == "research"
        else _ppt_chunks(snapshot)
    )

    for index, (title, lines) in enumerate(chunks, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        frame = _ppt_apply_content_theme(slide, snapshot["title"])
        frame.clear()

        # Variação sutil de acento visual por bloco.
        accent = PPT_BLUE if index % 3 == 1 else PPT_AMBER if index % 3 == 2 else PPT_RED
        accent_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.72),
            Inches(1.33),
            Inches(1.3),
            Inches(0.055),
        )
        accent_line.fill.solid()
        accent_line.fill.fore_color.rgb = accent
        accent_line.line.fill.background()

        for line_index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.level = 0
            paragraph.font.name = PROFESSIONAL_BRAND.font
            paragraph.font.size = PptPt(16)
            paragraph.font.color.rgb = PPT_TEXT
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.space_after = PptPt(7)

        _ppt_style_body(frame)

    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def render_section_export(project, section, export_format):
    snapshot = build_section_snapshot(project, section)
    renderers = {
        "html": render_section_html,
        "pdf": render_section_pdf,
        "docx": render_section_docx,
        "pptx": render_section_pptx,
    }
    try:
        renderer = renderers[export_format]
    except KeyError as exc:
        raise ValueError("Formato de exportação não suportado.") from exc
    return renderer(snapshot)
